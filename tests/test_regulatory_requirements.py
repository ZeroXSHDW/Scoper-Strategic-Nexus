import json
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from regulatory_requirements.cli import (
    ALLOWED_FRAMEWORKS,
    DocumentRecord,
    Source,
    build_index,
    clean_generated,
    clean_line,
    download_sources,
    export_artifacts,
    extract_requirements,
    load_sources,
    primary_source_file_path,
    quality_assessment,
    rel,
    response_required_for_requirement,
    sha256_bytes,
    should_keep_requirement,
    stable_id,
    validate_https_url,
)


class RegulatoryRequirementsTests(unittest.TestCase):
    def test_ci_and_docs_enforce_patch_hygiene(self):
        workflow = (Path(__file__).resolve().parents[1] / ".github/workflows/ci.yml").read_text(
            encoding="utf-8"
        )
        readme = (Path(__file__).resolve().parents[1] / "README.md").read_text(
            encoding="utf-8"
        )

        checkout_count = workflow.count("actions/checkout@")
        self.assertGreater(checkout_count, 0)
        self.assertEqual(checkout_count, workflow.count("git diff --check"))
        self.assertIn("runs-on: ubuntu-24.04", workflow)
        self.assertNotIn("runs-on: ubuntu-latest", workflow)
        self.assertNotIn("pip install --upgrade pip", workflow)
        self.assertNotIn("pip install --upgrade pip setuptools wheel", workflow)
        self.assertIn("pip install --disable-pip-version-check -e . pytest build", workflow)
        self.assertIn("git diff --check", readme)

    def test_manifest_loads_unique_sources(self):
        sources = load_sources()
        self.assertGreaterEqual(len(sources), 10)
        self.assertEqual(len({source.id for source in sources}), len(sources))
        self.assertTrue({source.framework for source in sources}.issubset(ALLOWED_FRAMEWORKS))
        self.assertIn("FRB_SUPERVISION", {source.framework for source in sources})
        self.assertIn("FFIEC", {source.framework for source in sources})

    def test_manifest_accepts_empty_sources_list(self):
        with tempfile.TemporaryDirectory() as tmp:
            manifest = Path(tmp) / "sources.json"
            manifest.write_text(
                json.dumps({"schema_version": 1, "sources": []}),
                encoding="utf-8",
            )
            self.assertEqual(load_sources(manifest), [])

    def test_manifest_rejects_missing_required_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            manifest = Path(tmp) / "sources.json"
            manifest.write_text(
                json.dumps(
                    {
                        "schema_version": 1,
                        "sources": [
                            {
                                "id": "incomplete",
                                "title": "Incomplete",
                                "framework": "FFIEC",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "missing"):
                load_sources(manifest)

    def test_primary_source_path_handles_absolute_and_relative(self):
        with tempfile.TemporaryDirectory() as tmp:
            absolute = Path(tmp) / "absolute.txt"
            absolute.write_text("x", encoding="utf-8")
            source = Source(
                id="path-fixture",
                title="Path fixture",
                authority="Fixture",
                framework="FFIEC",
                jurisdiction="US",
                source_url="https://example.com/source",
                download_url=None,
                local_path=str(absolute),
                source_kind="official_text",
                parser_profile="frb_ffiec",
                access="public_direct",
                notes="fixture",
            )
            self.assertEqual(
                primary_source_file_path(source).resolve(),
                absolute.resolve(),
            )
            # Absolute paths outside PROJECT_ROOT fall back to full path strings.
            self.assertEqual(rel(absolute), str(absolute.resolve()))

    def test_manifest_rejects_non_target_frameworks(self):
        with tempfile.TemporaryDirectory() as tmp:
            manifest = Path(tmp) / "sources.json"
            manifest.write_text(
                json.dumps(
                    {
                        "schema_version": 1,
                        "sources": [
                            {
                                "id": "unsupported-nist",
                                "title": "Unsupported Framework",
                                "authority": "NIST",
                                "framework": "NIST",
                                "jurisdiction": "United States",
                                "source_url": "https://example.com/nist.pdf",
                                "download_url": None,
                                "local_path": "nist.pdf",
                                "source_kind": "official_pdf",
                                "parser_profile": "frb_ffiec",
                                "access": "public_direct",
                                "notes": "Should not be allowed in this focused toolkit.",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "unsupported framework"):
                load_sources(manifest)

    def test_url_validation_rejects_non_https_and_private_hosts(self):
        with self.assertRaises(ValueError):
            validate_https_url("http://example.com/file.pdf")
        with self.assertRaises(ValueError):
            validate_https_url("https://127.0.0.1/file.pdf")

    def test_download_uses_local_fallback_when_online_refresh_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fallback = root / "fallback.txt"
            fallback.write_text("Management should maintain documented evidence.", encoding="utf-8")
            source = Source(
                id="fixture-fallback",
                title="Fixture fallback",
                authority="Fixture",
                framework="FFIEC",
                jurisdiction="US",
                source_url="https://127.0.0.1/source.pdf",
                download_url=None,
                local_path=str(root / "missing.pdf"),
                source_kind="official_pdf",
                parser_profile="frb_ffiec",
                access="public_direct",
                notes="fixture",
                fallback_local_paths=(str(fallback),),
            )
            results = download_sources([source], output_dir=root, refresh=True)
            self.assertEqual(results[0]["status"], "fallback_available")
            self.assertEqual(results[0]["freshness_status"], "check_failed")
            self.assertEqual(results[0]["local_path"], str(fallback.resolve()))

    def test_hash_and_stable_id_are_deterministic(self):
        self.assertEqual(sha256_bytes(b"abc"), sha256_bytes(b"abc"))
        self.assertEqual(stable_id("DORA", "a", "b"), stable_id("DORA", "a", "b"))
        self.assertNotEqual(stable_id("DORA", "a", "b"), stable_id("DORA", "a", "c"))

    def test_clean_line_strips_dora_publication_headers_and_ocr_artifacts(self):
        dirty = (
            "EN Official Journal of the European Union L 333/34 27.12.2022 "
            "7. Financial entities shall infor m management of an y op tion theret o "
            "and keep evidence av ailable."
        )
        cleaned = clean_line(dirty)
        self.assertTrue(cleaned.startswith("7. Financial entities shall inform"))
        self.assertIn("any option thereto", cleaned)
        self.assertIn("available", cleaned)
        self.assertNotIn("Official Journal", cleaned)

    def test_publication_only_noise_is_not_kept_as_requirement(self):
        self.assertFalse(should_keep_requirement("dora", "EN Official Journal of the European Union L 333/34 27.12.2022"))

    def test_quality_assessment_flags_source_and_content_review(self):
        row = {
            "source_hash": "",
            "extraction_confidence": 0.7,
            "requirement_text": "Where a Member State makes use of such op tion.",
            "citation": "General",
            "domain": "General regulatory obligation",
        }
        source = {"status": "failed", "access": "public_direct", "sha256": None}
        flag, note = quality_assessment(row, source)
        self.assertEqual(flag, "Source review")
        self.assertIn("Source unavailable", note)
        self.assertIn("Missing requirement source hash", note)

    def test_response_required_classifies_reference_only_definitions(self):
        row = {
            "framework": "DORA",
            "domain": "General regulatory obligation",
            "section_title": "Article 3",
            "normalized_obligation": "Definitions For the purposes of this Regulation, the following definitions shall apply:",
            "requirement_text": "Definitions For the purposes of this Regulation, the following definitions shall apply:",
        }
        self.assertEqual(response_required_for_requirement(row, {"status": "cached"}), "No")
        row["normalized_obligation"] = "Financial entities shall implement an ICT risk management framework."
        row["requirement_text"] = row["normalized_obligation"]
        row["domain"] = "ICT risk management"
        self.assertEqual(response_required_for_requirement(row, {"status": "cached"}), "Yes")

    def test_clause_extraction_for_dora_article_text(self):
        source = Source(
            id="fixture-dora",
            title="Fixture DORA",
            authority="Fixture Authority",
            framework="DORA",
            jurisdiction="EU",
            source_url="https://example.com/dora",
            download_url=None,
            local_path="fixture.txt",
            source_kind="official_text",
            parser_profile="dora",
            access="public_direct",
            notes="fixture",
        )
        document = DocumentRecord(
            document_id="DOC-fixture",
            source_id=source.id,
            title=source.title,
            path="fixture.txt",
            doc_type="text",
            sha256="abc123",
            page_count=None,
            parser_profile="dora",
            extracted_at="2026-05-20T00:00:00+00:00",
        )
        text = """
        Article 6
        ICT risk management framework
        Financial entities shall have a sound and comprehensive ICT risk management framework.
        Financial entities shall monitor the effectiveness of the implementation of their digital operational resilience strategy.
        """
        requirements = extract_requirements(source, document, text)
        self.assertGreaterEqual(len(requirements), 2)
        self.assertTrue(all(req.framework == "DORA" for req in requirements))
        self.assertTrue(any(req.domain == "ICT risk management" for req in requirements))

    def test_build_index_from_text_fixture(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_file = root / "fixture.txt"
            source_file.write_text(
                "I.A Security Culture\n"
                "Management should maintain an information security program and report assurance results.\n",
                encoding="utf-8",
            )
            manifest_source = Source(
                id="fixture-ffiec",
                title="Fixture FFIEC",
                authority="FFIEC",
                framework="FFIEC",
                jurisdiction="US",
                source_url="https://example.com/ffiec",
                download_url=None,
                local_path=str(source_file),
                source_kind="official_text",
                parser_profile="frb_ffiec",
                access="public_direct",
                notes="fixture",
            )
            db_path = root / "index.db"
            summary = build_index([manifest_source], db_path=db_path, output_dir=root)
            self.assertEqual(summary["documents"], 1)
            self.assertGreaterEqual(summary["requirements"], 1)
            conn = sqlite3.connect(db_path)
            try:
                count = conn.execute("SELECT COUNT(*) FROM requirements").fetchone()[0]
            finally:
                conn.close()
            self.assertEqual(count, summary["requirements"])

    def test_failed_index_rebuild_preserves_existing_database(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            db_path = root / "index.db"
            with sqlite3.connect(db_path) as conn:
                conn.execute("CREATE TABLE sentinel (value TEXT NOT NULL)")
                conn.execute("INSERT INTO sentinel(value) VALUES ('known-good')")
                conn.commit()

            source = Source(
                id="fixture-failure",
                title="Fixture failure",
                authority="Fixture",
                framework="FFIEC",
                jurisdiction="US",
                source_url="https://example.com/ffiec",
                download_url=None,
                local_path=str(root / "fixture.txt"),
                source_kind="official_text",
                parser_profile="frb_ffiec",
                access="public_direct",
                notes="fixture",
            )
            (root / "fixture.txt").write_text("fixture", encoding="utf-8")

            with patch(
                "regulatory_requirements.cli.document_text_and_record",
                side_effect=RuntimeError("simulated parser failure"),
            ):
                with self.assertRaisesRegex(RuntimeError, "simulated parser failure"):
                    build_index([source], db_path=db_path, output_dir=root)

            with sqlite3.connect(db_path) as conn:
                self.assertEqual(conn.execute("SELECT value FROM sentinel").fetchone()[0], "known-good")

    def test_export_artifacts_creates_vendor_ready_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_file = root / "fixture.txt"
            source_file.write_text(
                "I.A Security Culture\n"
                "Management should maintain an information security program, document ownership, and provide evidence.\n"
                "Management must review security testing results and track remediation actions.\n",
                encoding="utf-8",
            )
            manifest_source = Source(
                id="fixture-ffiec",
                title="Fixture FFIEC",
                authority="FFIEC",
                framework="FFIEC",
                jurisdiction="US",
                source_url="https://example.com/ffiec",
                download_url=None,
                local_path=str(source_file),
                source_kind="official_text",
                parser_profile="frb_ffiec",
                access="public_direct",
                notes="fixture",
            )
            db_path = root / "index.db"
            build_index([manifest_source], db_path=db_path, output_dir=root)
            summary = export_artifacts(db_path=db_path, output_dir=root)
            xlsx_path = root / Path(summary["xlsx_path"]).name
            docx_path = root / Path(summary["docx_path"]).name
            self.assertTrue(docx_path.is_file())
            wb = load_workbook(xlsx_path)
            try:
                expected = {
                    "Dashboard",
                    "Requirements",
                    "Vendor Questions",
                    "Vendor Worklist",
                    "Vendor - US Bank",
                    "Evidence Checklist",
                    "QA Flags",
                    "Source Issues",
                }
                self.assertTrue(expected.issubset(set(wb.sheetnames)))
                vendor = wb["Vendor Questions"]
                headers = [cell.value for cell in vendor[1]]
                for header in [
                    "priority",
                    "response_required",
                    "applicability",
                    "implementation_status",
                    "evidence_status",
                    "gap_or_exception_notes",
                    "completion_status",
                ]:
                    self.assertIn(header, headers)
                self.assertGreaterEqual(len(vendor.data_validations.dataValidation), 5)
                self.assertEqual(wb["Requirements"].max_row - 1, summary["database_requirements"])
                self.assertGreaterEqual(wb["Vendor Worklist"].max_row, 2)
                self.assertGreaterEqual(wb["Evidence Checklist"].max_row, 2)
            finally:
                wb.close()
            conn = sqlite3.connect(db_path)
            try:
                artifact_rows = conn.execute(
                    "SELECT artifact_type, path FROM artifact_runs ORDER BY artifact_type"
                ).fetchall()
            finally:
                conn.close()
            self.assertEqual({row[0] for row in artifact_rows}, {"docx", "xlsx"})
            self.assertTrue(all(row[1] for row in artifact_rows))

            with zipfile.ZipFile(docx_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
            for required_text in [
                "Vendor and Test Profile",
                "Regulatory Applicability",
                "Regulatory Scoping Questionnaire",
                "Evidence Package Checklist",
                "Exception and Remediation Register",
                "Source Review Acknowledgement",
                "Regulatory Coverage Matrix",
                "Vendor Attestation and Sign-Off",
                "ATTESTATION_REQUIRED",
                f"MAPPED_REQUIREMENTS_TOTAL:{summary['database_requirements']}",
                "UNMAPPED_REQUIREMENTS:0",
            ]:
                self.assertIn(required_text, document_xml)
            self.assertIn("ICT risk management", document_xml)
            self.assertIn("ICT-01", document_xml)
            self.assertNotIn("Vendor Response / Gap Notes", document_xml)

    def test_clean_generated_keeps_latest_outputs_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            keep_xlsx = root / "regulatory_requirements_reference_latest.xlsx"
            keep_docx = root / "regulatory_vendor_questionnaire_template_latest.docx"
            remove_xlsx = root / "regulatory_requirements_reference_2026-05-21.xlsx"
            remove_docx = root / "regulatory_vendor_questionnaire_template_2026-05-21.docx"
            for path in [keep_xlsx, keep_docx, remove_xlsx, remove_docx, root / "index.db", root / "download-log.json"]:
                path.write_text("fixture", encoding="utf-8")
            render_dir = root / "docx_render_check"
            render_dir.mkdir()
            (render_dir / "page-1.png").write_text("fixture", encoding="utf-8")

            result = clean_generated(root)

            self.assertTrue(keep_xlsx.exists())
            self.assertTrue(keep_docx.exists())
            self.assertFalse(remove_xlsx.exists())
            self.assertFalse(remove_docx.exists())
            self.assertFalse((root / "index.db").exists())
            self.assertFalse((root / "download-log.json").exists())
            self.assertFalse(render_dir.exists())
            self.assertGreaterEqual(len(result["removed"]), 5)


if __name__ == "__main__":
    unittest.main()
