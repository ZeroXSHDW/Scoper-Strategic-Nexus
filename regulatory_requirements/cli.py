#!/usr/bin/env python3
"""Download, index, and export the regulatory scoping workbook and questionnaire.

The pipeline is intentionally self-contained and does not rely on any web
application runtime.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import ipaddress
import json
import re
import shutil
import socket
import sqlite3
import ssl
import sys
import textwrap
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - surfaced by dependency check.
    PdfReader = None  # type: ignore[assignment]

try:
    from lxml import html as lxml_html
except ImportError:  # pragma: no cover - fallback is intentionally simple.
    lxml_html = None  # type: ignore[assignment]

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover - surfaced by dependency check.
    Document = None  # type: ignore[assignment]
    WD_ORIENT = None  # type: ignore[assignment]
    WD_CELL_VERTICAL_ALIGNMENT = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    OxmlElement = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    Inches = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    RGBColor = None  # type: ignore[assignment]

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:  # pragma: no cover - surfaced by dependency check.
    Workbook = None  # type: ignore[assignment]
    load_workbook = None  # type: ignore[assignment]
    BarChart = Reference = None  # type: ignore[assignment]
    FormulaRule = None  # type: ignore[assignment]
    Alignment = Border = Font = PatternFill = Side = None  # type: ignore[assignment]
    get_column_letter = None  # type: ignore[assignment]
    DataValidation = None  # type: ignore[assignment]
    Table = TableStyleInfo = None  # type: ignore[assignment]


PACKAGE_DIR = Path(__file__).resolve().parent
REPO_ROOT = PACKAGE_DIR.parent
PROJECT_ROOT = REPO_ROOT / "Penetration Testing - Scoping"
WORKSPACE_ROOT = REPO_ROOT
DEFAULT_MANIFEST = PACKAGE_DIR / "sources.json"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "Generated" / "Regulatory"
DEFAULT_DB = DEFAULT_OUTPUT_DIR / "index.db"
MAX_DOWNLOAD_BYTES = 40 * 1024 * 1024
USER_AGENT = "ZeroDev-Regulatory-Requirements/1.0 (+https://zerodevllc.com)"
ALLOWED_FRAMEWORKS = {
    "DORA",
    "DORA_TLPT",
    "FRB_SUPERVISION",
    "FFIEC",
    "HKMA_ICAST",
    "PRA_CBEST",
    "SCOPE_TEMPLATE",
}
REGULATORY_FOCUS = (
    "DORA, FRB/FFIEC supervision, HKMA CFI/iCAST, and PRA/Bank of England CBEST."
)


@dataclass(frozen=True)
class Source:
    id: str
    title: str
    authority: str
    framework: str
    jurisdiction: str
    source_url: str
    download_url: str | None
    local_path: str
    source_kind: str
    parser_profile: str
    access: str
    notes: str
    fallback_local_paths: tuple[str, ...] = ()


@dataclass(frozen=True)
class DocumentRecord:
    document_id: str
    source_id: str
    title: str
    path: str
    doc_type: str
    sha256: str | None
    page_count: int | None
    parser_profile: str
    extracted_at: str


@dataclass(frozen=True)
class Requirement:
    requirement_id: str
    source_id: str
    document_id: str
    framework: str
    authority: str
    jurisdiction: str
    citation: str
    section_title: str
    domain: str
    requirement_text: str
    normalized_obligation: str
    evidence_expectation: str
    vendor_question: str
    response_type: str
    service_codes: str
    source_hash: str
    extraction_confidence: float
    created_at: str


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def today_slug() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_sources(manifest_path: Path = DEFAULT_MANIFEST) -> list[Source]:
    payload = read_json(manifest_path)
    raw_sources = payload.get("sources", [])
    if not isinstance(raw_sources, list):
        raise ValueError("Manifest 'sources' must be a list")
    sources: list[Source] = []
    required = {
        "id",
        "title",
        "authority",
        "framework",
        "jurisdiction",
        "source_url",
        "local_path",
        "source_kind",
        "parser_profile",
        "access",
        "notes",
    }
    seen: set[str] = set()
    for item in raw_sources:
        missing = sorted(required.difference(item))
        if missing:
            raise ValueError(f"Manifest source is missing {', '.join(missing)}")
        source_id = str(item["id"])
        if source_id in seen:
            raise ValueError(f"Duplicate source id: {source_id}")
        seen.add(source_id)
        framework = str(item["framework"])
        if framework not in ALLOWED_FRAMEWORKS:
            allowed = ", ".join(sorted(ALLOWED_FRAMEWORKS))
            raise ValueError(
                f"Source {source_id} uses unsupported framework {framework!r}. "
                f"This toolkit is intentionally limited to: {allowed}"
            )
        sources.append(
            Source(
                id=source_id,
                title=str(item["title"]),
                authority=str(item["authority"]),
                framework=framework,
                jurisdiction=str(item["jurisdiction"]),
                source_url=str(item["source_url"]),
                download_url=item.get("download_url"),
                local_path=str(item["local_path"]),
                source_kind=str(item["source_kind"]),
                parser_profile=str(item["parser_profile"]),
                access=str(item["access"]),
                notes=str(item["notes"]),
                fallback_local_paths=tuple(str(path) for path in item.get("fallback_local_paths", [])),
            )
        )
    return sources


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_id(prefix: str, *parts: str) -> str:
    digest = hashlib.sha1("||".join(parts).encode("utf-8")).hexdigest()[:12]
    cleaned = re.sub(r"[^A-Z0-9]+", "_", prefix.upper()).strip("_") or "REQ"
    return f"{cleaned}-{digest}"


def primary_source_file_path(source: Source) -> Path:
    raw = Path(source.local_path)
    if raw.is_absolute():
        return raw
    return (PROJECT_ROOT / raw).resolve()


def source_file_path(source: Source) -> Path:
    primary = primary_source_file_path(source)
    if primary.is_file():
        return primary
    for fallback in source.fallback_local_paths:
        raw = Path(fallback)
        candidate = raw if raw.is_absolute() else (PROJECT_ROOT / raw).resolve()
        if candidate.is_file():
            return candidate
    return primary


def source_file_is_available(source: Source) -> bool:
    return source_file_path(source).is_file()


def download_log_path(output_dir: Path = DEFAULT_OUTPUT_DIR) -> Path:
    return output_dir / "download-log.json"


def ensure_dependencies(command: str) -> None:
    missing: list[str] = []
    if command in {"index", "build-all"} and PdfReader is None:
        missing.append("pypdf")
    if command in {"export", "build-all"}:
        if Workbook is None:
            missing.append("openpyxl")
        if Document is None:
            missing.append("python-docx")
    if missing:
        raise RuntimeError(
            "Missing Python package(s): "
            + ", ".join(sorted(set(missing)))
            + ". Install the packages listed in requirements.txt."
        )


def is_private_or_reserved(host: str) -> bool:
    try:
        address = ipaddress.ip_address(host)
        return (
            address.is_private
            or address.is_loopback
            or address.is_link_local
            or address.is_multicast
            or address.is_reserved
            or address.is_unspecified
        )
    except ValueError:
        return False


def validate_https_url(value: str) -> None:
    parsed = urlparse(value)
    if parsed.scheme != "https":
        raise ValueError(f"Only HTTPS URLs are allowed for downloads: {value}")
    if parsed.username or parsed.password:
        raise ValueError("Download URLs must not include credentials")
    host = parsed.hostname
    if not host:
        raise ValueError(f"URL is missing a host: {value}")
    if is_private_or_reserved(host):
        raise ValueError(f"URL resolves to a private or reserved literal address: {value}")
    try:
        infos = socket.getaddrinfo(host, None, proto=socket.IPPROTO_TCP)
    except socket.gaierror as exc:
        raise ValueError(f"Unable to resolve host for {value}: {exc}") from exc
    for info in infos:
        address = info[4][0]
        if is_private_or_reserved(address):
            raise ValueError(f"Host resolves to a private or reserved address: {value}")


def fetch_bytes(url: str, timeout: int = 30) -> tuple[bytes, str, str]:
    validate_https_url(url)
    context = ssl.create_default_context()
    req = Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/pdf,text/html,application/xhtml+xml,application/octet-stream;q=0.8,*/*;q=0.5",
        },
    )
    with urlopen(req, timeout=timeout, context=context) as response:
        final_url = response.geturl()
        validate_https_url(final_url)
        content_type = response.headers.get("content-type", "application/octet-stream").split(";")[0].strip()
        chunks: list[bytes] = []
        total = 0
        while True:
            chunk = response.read(1024 * 1024)
            if not chunk:
                break
            total += len(chunk)
            if total > MAX_DOWNLOAD_BYTES:
                raise ValueError(f"Download exceeds {MAX_DOWNLOAD_BYTES} bytes")
            chunks.append(chunk)
    return b"".join(chunks), content_type, final_url


def discover_pdf_url(base_url: str, html_body: bytes) -> str | None:
    text = html_body.decode("utf-8", errors="replace")
    candidates: list[str] = []
    for href in re.findall(r"href=[\"']([^\"']+)[\"']", text, flags=re.IGNORECASE):
        absolute = urljoin(base_url, html.unescape(href))
        lowered = absolute.lower()
        if ".pdf" in lowered or "download" in lowered or "@@download" in lowered:
            candidates.append(absolute)
    for candidate in candidates:
        lowered = candidate.lower()
        if ".pdf" in lowered:
            return candidate
    return candidates[0] if candidates else None


def should_promote_discovered_pdf(source: Source, target: Path) -> bool:
    return target.suffix.lower() == ".pdf" or source.source_kind.endswith("_pdf")


def download_sources(
    sources: list[Source],
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    force: bool = False,
    offline: bool = False,
    refresh: bool = False,
) -> list[dict[str, Any]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    for source in sources:
        target = primary_source_file_path(source)
        previous_sha = sha256_file(target) if target.exists() else None
        result: dict[str, Any] = {
            "source_id": source.id,
            "title": source.title,
            "framework": source.framework,
            "source_kind": source.source_kind,
            "local_path": rel(target) if target.exists() else str(source.local_path),
            "status": "pending",
            "sha256": previous_sha,
            "previous_sha256": previous_sha,
            "size_bytes": target.stat().st_size if target.exists() else None,
            "content_type": None,
            "final_url": None,
            "latest_url": None,
            "latest_sha256": None,
            "discovered_pdf_url": None,
            "source_page_url": None,
            "source_page_sha256": None,
            "source_page_status": "unchecked",
            "source_page_checked_at": None,
            "freshness_status": "unchecked",
            "freshness_checked_at": None,
            "fetched_at": None,
            "message": "",
        }

        if source.source_kind == "supplied_pdf":
            result["status"] = "available" if target.is_file() else "missing"
            result["source_page_status"] = "workspace_supplied"
            result["freshness_status"] = "workspace_supplied"
            result["message"] = "Workspace-supplied file" if target.is_file() else "Workspace-supplied file missing"
            results.append(result)
            continue

        if offline:
            result["status"] = "missing" if not target.is_file() else "cached"
            result["source_page_status"] = "offline_skipped"
            result["freshness_status"] = "offline_skipped"
            result["message"] = "Offline mode; freshness check and download skipped"
            results.append(result)
            continue

        if target.is_file() and not force and not refresh:
            result["status"] = "cached"
            result["source_page_status"] = "not_checked"
            result["freshness_status"] = "not_checked"
            result["message"] = "Using existing local archive copy"
            results.append(result)
            continue

        url = source.download_url or source.source_url
        if refresh and source.download_url and source.source_url != source.download_url:
            try:
                page_data, page_type, page_final_url = fetch_bytes(source.source_url)
                result.update(
                    {
                        "source_page_url": page_final_url,
                        "source_page_sha256": sha256_bytes(page_data),
                        "source_page_status": "checked",
                        "source_page_checked_at": now_iso(),
                    }
                )
                if "html" in page_type.lower() or b"<html" in page_data[:5000].lower():
                    result["discovered_pdf_url"] = discover_pdf_url(page_final_url, page_data)
            except (HTTPError, URLError, TimeoutError, ValueError, ssl.SSLError, OSError) as exc:
                result.update(
                    {
                        "source_page_url": source.source_url,
                        "source_page_status": "check_failed",
                        "source_page_checked_at": now_iso(),
                        "message": f"Source page freshness check failed: {exc}",
                    }
                )
        try:
            data, content_type, final_url = fetch_bytes(url)
            looks_html = "html" in content_type.lower() or b"<html" in data[:5000].lower()
            discovered = None
            if looks_html and not data.startswith(b"%PDF-"):
                discovered = discover_pdf_url(final_url, data)
                result["discovered_pdf_url"] = discovered or result["discovered_pdf_url"]
                if discovered and should_promote_discovered_pdf(source, target):
                    try:
                        pdf_data, pdf_type, pdf_final_url = fetch_bytes(discovered)
                        if pdf_data.startswith(b"%PDF-"):
                            data, content_type, final_url = pdf_data, pdf_type, pdf_final_url
                    except (HTTPError, URLError, TimeoutError, ValueError, ssl.SSLError):
                        pass
            if target.suffix.lower() == ".pdf" and not data.startswith(b"%PDF-"):
                raise ValueError(f"Expected PDF but received {content_type or 'unknown content type'}")

            latest_sha = sha256_bytes(data)
            status = "downloaded"
            message = "Downloaded official source"
            if previous_sha and latest_sha == previous_sha and refresh and not force:
                status = "current"
                message = "Official source checked; local archive is current"
            else:
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(data)
                if previous_sha:
                    status = "updated"
                    message = "Official source changed; local archive updated"
            page_note = ""
            if result.get("source_page_status") == "check_failed":
                page_note = " Source landing page could not be checked; archived document was still refreshed."
            result.update(
                {
                    "status": status,
                    "sha256": latest_sha,
                    "latest_sha256": latest_sha,
                    "size_bytes": len(data),
                    "content_type": content_type,
                    "final_url": final_url,
                    "latest_url": final_url,
                    "freshness_status": "current" if status == "current" else "updated" if status == "updated" else "downloaded",
                    "freshness_checked_at": now_iso(),
                    "fetched_at": now_iso(),
                    "local_path": rel(target),
                    "message": f"{message}{page_note}",
                }
            )
        except (HTTPError, URLError, TimeoutError, ValueError, ssl.SSLError, OSError) as exc:
            fallback_path = source_file_path(source)
            if fallback_path.is_file():
                fallback_hash = sha256_file(fallback_path)
                result.update(
                    {
                        "status": "fallback_available" if fallback_path != target else "cached_refresh_failed",
                        "sha256": fallback_hash,
                        "size_bytes": fallback_path.stat().st_size,
                        "local_path": rel(fallback_path),
                        "freshness_status": "check_failed",
                        "freshness_checked_at": now_iso(),
                        "message": f"{exc}; using local archive fallback for indexing",
                    }
                )
            else:
                result["status"] = "failed"
                result["freshness_status"] = "check_failed"
                result["freshness_checked_at"] = now_iso()
                result["message"] = str(exc)
        results.append(result)

    log_path = download_log_path(output_dir)
    log_path.write_text(json.dumps({"generated_at": now_iso(), "results": results}, indent=2), encoding="utf-8")
    return results


def create_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        PRAGMA foreign_keys = ON;

        DROP TABLE IF EXISTS requirement_mappings;
        DROP TABLE IF EXISTS requirements;
        DROP TABLE IF EXISTS documents;
        DROP TABLE IF EXISTS sources;
        DROP TABLE IF EXISTS artifact_runs;

        CREATE TABLE sources (
            source_id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            authority TEXT NOT NULL,
            framework TEXT NOT NULL,
            jurisdiction TEXT NOT NULL,
            source_url TEXT NOT NULL,
            download_url TEXT,
            local_path TEXT NOT NULL,
            source_kind TEXT NOT NULL,
            parser_profile TEXT NOT NULL,
            access TEXT NOT NULL,
            status TEXT NOT NULL,
            sha256 TEXT,
            size_bytes INTEGER,
            content_type TEXT,
            final_url TEXT,
            fetched_at TEXT,
            latest_url TEXT,
            latest_sha256 TEXT,
            source_page_url TEXT,
            source_page_sha256 TEXT,
            source_page_status TEXT,
            source_page_checked_at TEXT,
            freshness_status TEXT,
            freshness_checked_at TEXT,
            notes TEXT
        );

        CREATE TABLE documents (
            document_id TEXT PRIMARY KEY,
            source_id TEXT NOT NULL REFERENCES sources(source_id) ON DELETE CASCADE,
            title TEXT NOT NULL,
            path TEXT NOT NULL,
            doc_type TEXT NOT NULL,
            sha256 TEXT,
            page_count INTEGER,
            parser_profile TEXT NOT NULL,
            extracted_at TEXT NOT NULL
        );

        CREATE TABLE requirements (
            requirement_id TEXT PRIMARY KEY,
            source_id TEXT NOT NULL REFERENCES sources(source_id) ON DELETE CASCADE,
            document_id TEXT NOT NULL REFERENCES documents(document_id) ON DELETE CASCADE,
            framework TEXT NOT NULL,
            authority TEXT NOT NULL,
            jurisdiction TEXT NOT NULL,
            citation TEXT NOT NULL,
            section_title TEXT NOT NULL,
            domain TEXT NOT NULL,
            requirement_text TEXT NOT NULL,
            normalized_obligation TEXT NOT NULL,
            evidence_expectation TEXT NOT NULL,
            vendor_question TEXT NOT NULL,
            response_type TEXT NOT NULL,
            service_codes TEXT NOT NULL,
            source_hash TEXT NOT NULL,
            extraction_confidence REAL NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(framework, citation, requirement_text, source_hash)
        );

        CREATE TABLE requirement_mappings (
            mapping_id TEXT PRIMARY KEY,
            requirement_id TEXT NOT NULL REFERENCES requirements(requirement_id) ON DELETE CASCADE,
            mapping_type TEXT NOT NULL,
            mapping_value TEXT NOT NULL,
            rationale TEXT NOT NULL
        );

        CREATE TABLE artifact_runs (
            run_id TEXT PRIMARY KEY,
            artifact_type TEXT NOT NULL,
            path TEXT NOT NULL,
            sha256 TEXT,
            generated_at TEXT NOT NULL,
            requirement_count INTEGER NOT NULL,
            source_count INTEGER NOT NULL,
            notes TEXT
        );

        CREATE INDEX idx_requirements_framework ON requirements(framework);
        CREATE INDEX idx_requirements_domain ON requirements(domain);
        CREATE INDEX idx_requirements_citation ON requirements(citation);
        """
    )


def load_download_results(output_dir: Path = DEFAULT_OUTPUT_DIR) -> dict[str, dict[str, Any]]:
    log_path = download_log_path(output_dir)
    if not log_path.is_file():
        return {}
    payload = read_json(log_path)
    return {str(item["source_id"]): item for item in payload.get("results", [])}


def insert_source(conn: sqlite3.Connection, source: Source, download_result: dict[str, Any] | None) -> None:
    path = source_file_path(source)
    exists = path.is_file()
    file_hash = sha256_file(path) if exists else None
    status = "available" if exists else "missing"
    size = path.stat().st_size if exists else None
    content_type = None
    final_url = None
    fetched_at = None
    latest_url = None
    latest_sha256 = None
    source_page_url = None
    source_page_sha256 = None
    source_page_status = "not_checked"
    source_page_checked_at = None
    freshness_status = "not_checked"
    freshness_checked_at = None
    if download_result:
        status = str(download_result.get("status") or status)
        if exists and status in {"failed", "missing"} and path != primary_source_file_path(source):
            status = "fallback_available"
        elif exists and status == "failed":
            status = "cached_refresh_failed"
        content_type = download_result.get("content_type")
        final_url = download_result.get("final_url")
        fetched_at = download_result.get("fetched_at")
        latest_url = download_result.get("latest_url") or final_url
        latest_sha256 = download_result.get("latest_sha256")
        source_page_url = download_result.get("source_page_url")
        source_page_sha256 = download_result.get("source_page_sha256")
        source_page_status = str(download_result.get("source_page_status") or source_page_status)
        source_page_checked_at = download_result.get("source_page_checked_at")
        freshness_status = str(download_result.get("freshness_status") or freshness_status)
        freshness_checked_at = download_result.get("freshness_checked_at")
        file_hash = file_hash or download_result.get("sha256")
        size = size or download_result.get("size_bytes")
    conn.execute(
        """
        INSERT INTO sources (
            source_id, title, authority, framework, jurisdiction, source_url,
            download_url, local_path, source_kind, parser_profile, access,
            status, sha256, size_bytes, content_type, final_url, fetched_at,
            latest_url, latest_sha256, source_page_url, source_page_sha256,
            source_page_status, source_page_checked_at, freshness_status,
            freshness_checked_at, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            source.id,
            source.title,
            source.authority,
            source.framework,
            source.jurisdiction,
            source.source_url,
            source.download_url,
            rel(path) if exists else source.local_path,
            source.source_kind,
            source.parser_profile,
            source.access,
            status,
            file_hash,
            size,
            content_type,
            final_url,
            fetched_at,
            latest_url,
            latest_sha256,
            source_page_url,
            source_page_sha256,
            source_page_status,
            source_page_checked_at,
            freshness_status,
            freshness_checked_at,
            source.notes,
        ),
    )


def extract_pdf_text(path: Path) -> tuple[str, int]:
    if PdfReader is None:
        raise RuntimeError("pypdf is required to extract PDF text")
    reader = PdfReader(str(path))
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        parts.append(f"\n[[PAGE {index}]]\n{text}")
    return "\n".join(parts), len(reader.pages)


def extract_html_text(path: Path) -> str:
    raw = path.read_bytes()
    if lxml_html is not None:
        doc = lxml_html.fromstring(raw)
        for bad in doc.xpath("//script|//style|//noscript"):
            bad.drop_tree()
        text = doc.text_content()
    else:
        body = raw.decode("utf-8", errors="replace")
        body = re.sub(r"(?is)<(script|style).*?</\1>", " ", body)
        text = re.sub(r"(?s)<[^>]+>", " ", body)
    return html.unescape(text)


def clean_line(value: str) -> str:
    value = value.replace("\u00a0", " ")
    value = value.replace("\u2022", "-")
    value = value.replace("\u2013", "-").replace("\u2014", "-")
    value = value.replace("\u2018", "'").replace("\u2019", "'")
    value = value.replace("\u201c", '"').replace("\u201d", '"')
    value = re.sub(r"\s+", " ", value)
    replacements = {
        "Ar ticle": "Article",
        "ar ticle": "article",
        "Reg ulation": "Regulation",
        "reg ulation": "regulation",
        "financia l": "financial",
        "Financia l": "Financial",
        "r isk": "risk",
        "R isk": "Risk",
        "framewor k": "framework",
        "Framewor k": "Framework",
        "compet ent": "competent",
        "Compet ent": "Competent",
        "oper ational": "operational",
        "Oper ational": "Operational",
        "infor mation": "information",
        "In for mation": "Information",
        "ser vices": "services",
        "Ser vices": "Services",
        "ser vice": "service",
        "Ser vice": "Service",
        "sect or": "sector",
        "Sect or": "Sector",
        "thir d": "third",
        "Thir d": "Third",
        "par ty": "party",
        "Par ty": "Party",
        "f or": "for",
        "F or": "For",
        "f ollowing": "following",
        "f ollowi ng": "following",
        "f inancial": "financial",
        "f inancia l": "financial",
        "F inancial": "Financial",
        "r ules": "rules",
        "R ules": "Rules",
        "refer red": "referred",
        "Refer red": "Referred",
        "Member Stat e": "Member State",
        "Member Stat es": "Member States",
        "inf or m": "inform",
        "Inf or m": "Inform",
        "c hanges": "changes",
        "C hanges": "Changes",
        "pur poses": "purposes",
        "Pur poses": "Purposes",
        "ma y": "may",
        "Ma y": "May",
        "excl ude": "exclude",
        "Excl ude": "Exclude",
        "w ork": "work",
        "W ork": "Work",
        "hav e": "have",
        "Hav e": "Have",
        "cr itical": "critical",
        "Cr itical": "Critical",
        "U nion": "Union",
        "Uni on": "Union",
        "Offi cial Jour nal": "Official Journal",
        "Def initions": "Definitions",
        "op tion": "option",
        "op tions": "options",
        "infor m": "inform",
        "Infor mation": "Information",
        "infor mation": "information",
        "infor med": "informed",
        "av ailable": "available",
        "av ailability": "availability",
        "co vered": "covered",
        "professiona l": "professional",
        "exc hang e": "exchange",
        "author ities": "authorities",
        "author ity": "authority",
        "designat ed": "designated",
        "excep t": "except",
        "vir tue": "virtue",
        "theret o": "thereto",
        "an y": "any",
        "ha ve": "have",
        "ha ving": "having",
        "effect iveness": "effectiveness",
        "eff ectiveness": "effectiveness",
        "strateg y": "strategy",
        "secur ity": "security",
        "integrit y": "integrity",
        "confidentialit y": "confidentiality",
        "authenticit y": "authenticity",
        "T esting": "Testing",
        "test ing": "testing",
        "manag ement": "management",
        "Manag ement": "Management",
        "repor te d": "reported",
        "repor t": "report",
        "relate d": "related",
        "maj or": "major",
        "stak eholders": "stakeholders",
        "interconne cted": "interconnected",
        "inve stment": "investment",
        "Simplif ied": "Simplified",
        "fi r ms": "firms",
        "pa yment": "payment",
        "exe mp te d": "exempted",
        "comm unication": "communication",
        "fi nancial": "financial",
        "Fi nancial": "Financial",
        "im por tant": "important",
        "Im por tant": "Important",
        "im plement": "implement",
        "im plemented": "implemented",
        "im plementation": "implementation",
        "Propor tionality": "Proportionality",
        "propor tionality": "proportionality",
        "pr inciple": "principle",
        "pr inciples": "principles",
        "Chapt er": "Chapter",
        "Chapt ers": "Chapters",
        "maint enance": "maintenance",
        "relat ed": "related",
        "ICT -related": "ICT-related",
        "ICT -relat ed": "ICT-related",
        "microent er pr ises": "microenterprises",
        "disr upti ons": "disruptions",
        "o verall": "overall",
        "o ver": "over",
        "rest oration": "restoration",
        "under take n": "undertaken",
        "under tak e": "undertake",
        "per iodically": "periodically",
        "diffe rent": "different",
        "scenar ios": "scenarios",
        "reco ver y": "recovery",
        "extern al": "external",
        "cr isis": "crisis",
        "notific ation": "notification",
        "aggregat ed": "aggregated",
        "provid ed": "provided",
        "pro viders": "providers",
        "pay ment": "payment",
        "insti tutions": "institutions",
        "inst itutions": "institutions",
        "tec hnical": "technical",
        "te chnical": "technical",
        "f ollow-up": "follow-up",
        "national la w": "national law",
        "CSIRT s": "CSIRTs",
        "fi ndings": "findings",
        "fi nd": "find",
        "term inate": "terminate",
        "for m": "form",
        "appropr iate": "appropriate",
        "appropr iat e": "appropriate",
        "manage ment": "management",
        "ref er red": "referred",
        "refe r red": "referred",
        "af ter": "after",
        "summar y": "summary",
        "attes tation": "attestation",
        "imp osing": "imposing",
        "im posed": "imposed",
        "impos ed": "imposed",
        "im pose": "impose",
        "im posing": "imposing",
        "compl iance": "compliance",
        "achi eved": "achieved",
        "per iodic": "periodic",
        "per iod": "period",
        "r ight": "right",
        "Ex ercise": "Exercise",
        "ex ercise": "exercise",
        "po w er": "power",
        "po wers": "powers",
        "administ rativ e": "administrative",
        "administrativ e": "administrative",
        "Com petent": "Competent",
        "com petent": "competent",
        "exe rcise": "exercise",
        "leg al": "legal",
        "framework s": "frameworks",
        "f ollo ws": "follows",
        "dela y": "delay",
        "ag ainst": "against",
        "af te r": "after",
        "oppor tunity": "opportunity",
        "whic h": "which",
        "provid er": "provider",
        "provid e": "provide",
        "provid ed": "provided",
        "author ised": "authorised",
        "Committ ee": "Committee",
        "ever y": "every",
        "ye ars": "years",
        "ye arly": "yearly",
        "conf idential": "confidential",
        "P arliament": "Parliament",
        "countr ies": "countries",
        "f ocusing": "focusing",
        "ev olution": "evolution",
        "im pact": "impact",
        "im pacts": "impacts",
        "chang es": "changes",
        "cor rective": "corrective",
        "deplo ying": "deploying",
        "deplo yment": "deployment",
        "strate gies": "strategies",
        "too ls": "tools",
        "complet e": "complete",
        "update d": "updated",
        "im proved": "improved",
        "im provements": "improvements",
        "der ived": "derived",
        "implemen tation": "implementation",
        "submitt ed": "submitted",
        "mech anisms": "mechanisms",
        "mechani sms": "mechanisms",
        "det ect": "detect",
        "prev ention": "prevention",
        "system s": "systems",
        "Lear ning": "Learning",
        "lear ning": "learning",
        "ev olving": "evolving",
        "par ticular": "particular",
        "cyber -attacks": "cyber-attacks",
        "disr upts": "disrupts",
        "disr upti on": "disruption",
        "mitig ate": "mitigate",
        "ICT - related": "ICT-related",
        "taske d": "tasked",
        "pur pose": "purpose",
        "signif icant": "significant",
        "consist ent": "consistent",
        "consiste nt": "consistent",
        "integrat ed": "integrated",
        "monitor ing": "monitoring",
        "monitori ng": "monitoring",
        "monit or ing": "monitoring",
        "documente d": "documented",
        "wa ys": "ways",
        "f acilitate": "facilitate",
        "f lo w": "flow",
        "under pin": "underpin",
        "super visor y": "supervisory",
        "converg ence": "convergence",
        "anony mised": "anonymised",
        "ma jor": "major",
        "manag e": "manage",
        "manag ed": "managed",
        "eff ective": "effective",
        "ar range ments": "arrangements",
        "intern al": "internal",
        "go ver ned": "governed",
        "go ver n": "govern",
        "terr itory": "territory",
        "car ried": "carried",
        "car ry": "carry",
        "seri ously": "seriously",
        "jeopar dise": "jeopardise",
        "jeopar dize": "jeopardize",
        "propor tionate": "proportionate",
        "par ties": "parties",
        "par ty": "party",
        "in volved": "involved",
        "appro aches": "approaches",
        "recoveri ng": "recovering",
        "necessar y": "necessary",
        "c hecks": "checks",
        "integr ity": "integrity",
        "per formed": "performed",
        "reconstr ucting": "reconstructing",
        "ext er nal": "external",
        "stakeho lders": "stakeholders",
        "implementa tion": "implementation",
        "Regula tion": "Regulation",
        "regula tion": "regulation",
        "regulator y": "regulatory",
        "regardi ng": "regarding",
        "duri ng": "during",
        "enteri ng": "entering",
        "ar rangement": "arrangement",
        "y early": "yearly",
        "aggreg ated": "aggregated",
        "For um": "Forum",
        "for um": "forum",
        "countr y": "country",
        "concer ned": "concerned",
        "inve stigations": "investigations",
        "terr it or y": "territory",
        "jur isdiction": "jurisdiction",
        "suppor ting": "supporting",
        "compreh ensive": "comprehensive",
        "benc hmark s": "benchmarks",
        "adopt ed": "adopted",
        "ESA s": "ESAs",
        "exper ts": "experts",
        "select ed": "selected",
        "o versight": "oversight",
        "f oster": "foster",
        "mitig ants": "mitigants",
        "transfer s": "transfers",
        "draf t": "draft",
        "exercisi ng": "exercising",
        "exer cising": "exercising",
        "additiona l": "additional",
        "specifica tion": "specification",
        "compl aints": "complaints",
        "withthe": "with the",
        "Euro pean": "European",
        "EUR OPEAN": "EUROPEAN",
        "P ARLIAMENT": "PARLIAMENT",
        "CO UNCIL": "COUNCIL",
        "REGUL A TION": "REGULATION",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = re.sub(
        r"^EN\s+Official\s+Journal\s+of\s+the\s+European\s+Union\s+"
        r"(?:(?:L|C)\s+\d+/\d+\s+)?"
        r"(?:\d{1,2}\.\d{1,2}\.\d{4}\s+)?"
        r"(?:(?:L|C)\s+\d+/\d+\s+)?",
        "",
        value,
        flags=re.IGNORECASE,
    )
    value = re.sub(r"^EN\s+(?=\d+\.)", "", value)
    value = re.sub(r"\s+([,.;:])", r"\1", value)
    value = re.sub(r"\(\s+", "(", value)
    value = re.sub(r"\s+\)", ")", value)
    value = re.sub(r"\s{2,}", " ", value)
    return value.strip()


def clean_text(value: str) -> str:
    lines = [clean_line(line) for line in value.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)


def is_noise(line: str) -> bool:
    lowered = line.lower()
    if not line:
        return True
    if "official journal of the european union" in lowered and len(line) < 120:
        return True
    if "strictly confidential" in lowered:
        return True
    if "nexus" in lowered and "page" in lowered:
        return True
    if lowered.startswith("page ") and len(line) < 30:
        return True
    if re.search(r"\bl\s+\d+/\d+\s+\d{1,2}\.\d{1,2}\.\d{4}\b", lowered):
        return True
    if "cookie" in lowered and "privacy" in lowered and len(line) < 160:
        return True
    return False


def is_heading(line: str) -> bool:
    if len(line) > 140:
        return False
    patterns = [
        r"^Article\s+\d+[A-Za-z]?\b",
        r"^CHAPTER\s+[IVXLC]+",
        r"^SECTION\s+\d+",
        r"^ANNEX\s+[A-ZIVXLC]+",
        r"^\d+(?:\.\d+){0,3}\s+[-A-Za-z]",
        r"^[IVXLC]+\.[A-Z](?:\.\d+)?\s+[-A-Za-z]",
        r"^Phase\s+\d+\b",
        r"^SEC\s+\d+\b",
        r"^[A-Z][A-Za-z /&-]+:$",
    ]
    if any(re.search(pattern, line) for pattern in patterns):
        return True
    words = line.split()
    return 2 <= len(words) <= 10 and line.isupper() and not re.search(r"[.;]$", line)


OBLIGATION_PATTERN = re.compile(
    r"\b("
    r"shall|must|required|requires|requirement|should|expected|expectation|need to|"
    r"ensure|establish|maintain|implement|document|approve|review|monitor|report|"
    r"conduct|perform|test|assess|identify|classify|manage|validate|provide|retain|"
    r"evidence|remediation|scope|approval|authori[sz]ation"
    r")\b",
    re.IGNORECASE,
)


def split_sentences(value: str) -> list[str]:
    value = re.sub(r"\s+", " ", value).strip()
    if len(value) <= 420:
        return [value] if value else []
    pieces = re.split(r"(?<=[.;:])\s+(?=[A-Z0-9(\[])|(?<=\.)\s+-\s+", value)
    out: list[str] = []
    current = ""
    for piece in pieces:
        piece = piece.strip()
        if not piece:
            continue
        if len(current) + len(piece) < 360:
            current = f"{current} {piece}".strip()
        else:
            if current:
                out.append(current)
            current = piece
    if current:
        out.append(current)
    return out


def section_blocks(text: str, profile: str) -> Iterable[tuple[str, str, str]]:
    citation = "General"
    section = "Overview"
    paragraph = ""
    skip_uncited_dora = profile in {"dora", "dora_tlpt"}
    for raw_line in clean_text(text).splitlines():
        line = clean_line(raw_line)
        if is_noise(line):
            continue
        page_match = re.match(r"^\[\[PAGE\s+(\d+)\]\]$", line)
        if page_match:
            continue
        if is_heading(line):
            if paragraph:
                if not (skip_uncited_dora and citation == "General"):
                    yield citation, section, paragraph
                paragraph = ""
            section = line
            article = re.match(r"^(Article\s+\d+[A-Za-z]?)\b(.*)$", line)
            numbered = re.match(r"^(\d+(?:\.\d+){0,3})\s+(.+)$", line)
            phase = re.match(r"^(Phase\s+\d+)\b(.*)$", line)
            if article:
                citation = article.group(1)
                if article.group(2).strip():
                    section = line
            elif phase:
                citation = phase.group(1)
            elif numbered and profile != "dora":
                citation = numbered.group(1)
            continue
        if profile == "scope_spec" and "Client to complete" in line:
            if paragraph:
                if not (skip_uncited_dora and citation == "General"):
                    yield citation, section, paragraph
                paragraph = ""
            yield citation, section, line
            continue
        if line.startswith(("-", "*")) or re.match(r"^\(?[a-z0-9ivx]{1,4}\)\s+", line, flags=re.I):
            if paragraph:
                if not (skip_uncited_dora and citation == "General"):
                    yield citation, section, paragraph
                paragraph = ""
            if not (skip_uncited_dora and citation == "General"):
                yield citation, section, line
            continue
        paragraph = f"{paragraph} {line}".strip()
        if len(paragraph) > 650 or re.search(r"[.;:]$", line):
            if not (skip_uncited_dora and citation == "General"):
                yield citation, section, paragraph
            paragraph = ""
    if paragraph:
        if not (skip_uncited_dora and citation == "General"):
            yield citation, section, paragraph


def profile_confidence(profile: str, text: str) -> float:
    base = {
        "dora": 0.88,
        "dora_tlpt": 0.87,
        "frb_ffiec": 0.84,
        "hkma_icast": 0.84,
        "pra_cbest": 0.88,
        "scope_spec": 0.78,
    }.get(profile, 0.7)
    if re.search(r"\bshall\b|\bmust\b|\brequired\b", text, flags=re.I):
        base += 0.04
    if len(text) < 60:
        base -= 0.08
    return max(0.55, min(0.98, base))


def infer_domain(profile: str, text: str, section: str) -> str:
    haystack = f"{section} {text}".lower()
    mapping = [
        ("ICT risk management", ["ict risk", "risk management", "information security program", "security program"]),
        ("Third-party risk", ["third-party", "third party", "outsourc", "provider", "vendor", "supply chain"]),
        ("Operational resilience", ["resilience", "business continuity", "recovery", "important business service", "critical function"]),
        ("Threat-led testing", ["tlpt", "threat-led", "red team", "intelligence-led", "icast", "cbest", "tiber"]),
        ("Penetration testing", ["penetration", "security testing", "assurance and testing", "vulnerability"]),
        ("Incident response", ["incident", "response", "reporting", "notification", "cyber threat"]),
        ("Governance and approvals", ["board", "management body", "approval", "accountability", "sign-off"]),
        ("Evidence and reporting", ["evidence", "report", "documentation", "retain", "record"]),
        ("Scope and assets", ["asset", "scope", "system", "critical or important", "inventory", "register"]),
        ("Tester qualification", ["crest", "certification", "qualification", "independence", "tester"]),
    ]
    for domain, needles in mapping:
        if any(needle in haystack for needle in needles):
            return domain
    if profile == "scope_spec":
        return "Vendor intake and scope completion"
    return "General regulatory obligation"


def infer_response_type(text: str, domain: str) -> str:
    lowered = text.lower()
    if "date" in lowered or "timeline" in lowered or "frequency" in lowered:
        return "date_or_frequency"
    if "evidence" in lowered or "report" in lowered or "document" in lowered or "provide" in lowered:
        return "upload_and_narrative"
    if "yes" in lowered and "no" in lowered:
        return "yes_no"
    if domain in {"Scope and assets", "Third-party risk"}:
        return "table"
    return "narrative"


def infer_service_codes(text: str, domain: str) -> str:
    haystack = f"{domain} {text}".lower()
    codes: list[str] = []
    if "external" in haystack or "internet-facing" in haystack or "perimeter" in haystack:
        codes.extend(["PT-NET-EXT", "VA-NET-EXT"])
    if "internal" in haystack or "network" in haystack:
        codes.extend(["PT-NET-INT", "VA-NET-INT"])
    if "application" in haystack or "web" in haystack:
        codes.append("PT-APP-WEB")
    if "api" in haystack:
        codes.append("PT-APP-API")
    if "cloud" in haystack or "ict third-party" in haystack:
        codes.append("PT-CLD-INF")
    if "segmentation" in haystack:
        codes.append("PT-NET-SEG")
    if "tlpt" in haystack or "red team" in haystack or "threat-led" in haystack or "cbest" in haystack or "icast" in haystack:
        codes.append("PT-ADV-RT")
    if "incident" in haystack or "recovery" in haystack or "continuity" in haystack:
        codes.append("IR-TABLETOP")
    if not codes:
        codes.append("VDD-REG")
    return ", ".join(dict.fromkeys(codes))


def evidence_expectation(profile: str, domain: str, text: str) -> str:
    lowered = text.lower()
    if "tlpt" in lowered or "cbest" in lowered or "icast" in lowered:
        return "Threat intelligence report, approved test plan, execution report, detection/response assessment, remediation plan, and regulator-ready closure evidence."
    if "third" in lowered or "provider" in lowered or "outsourc" in lowered:
        return "Provider register, contract/control evidence, approval to test, due-diligence record, and ongoing monitoring evidence."
    if "incident" in lowered or "report" in lowered:
        return "Incident response plan, notification workflow, test records, lessons learned, and remediation evidence."
    if "asset" in lowered or "scope" in lowered or "inventory" in lowered:
        return "Current asset inventory, criticality mapping, ownership record, and approved in-scope/out-of-scope register."
    if "business continuity" in lowered or "recovery" in lowered:
        return "BC/DR plans, test calendar, scenario results, RTO/RPO evidence, and action tracker."
    if domain == "Vendor intake and scope completion":
        return "Completed vendor response, owner approval, supporting attachment, and change-control reference where applicable."
    return "Policy/procedure evidence, control owner attestation, testing output, finding tracker, and remediation evidence."


def normalized_obligation(text: str) -> str:
    value = re.sub(r"^\(?[a-z0-9ivx]{1,4}\)\s+", "", text, flags=re.I)
    value = re.sub(r"^[-*]\s+", "", value)
    value = value.strip()
    if len(value) > 260:
        value = value[:257].rsplit(" ", 1)[0] + "..."
    return value


def vendor_question(framework: str, citation: str, obligation: str, domain: str) -> str:
    return (
        f"Describe how your organization satisfies {framework} {citation} for {domain}. "
        f"Address this obligation: {obligation} Provide current evidence or a remediation plan."
    )


def should_keep_requirement(profile: str, text: str) -> bool:
    clean = normalized_obligation(text)
    if len(clean) < 35 or len(clean) > 1400:
        return False
    lowered = clean.lower()
    if any(
        noise in lowered
        for noise in [
            "table of contents",
            "copyright",
            "breadcrumb",
            "last revision date",
            "official journal of the european union",
            "page intentionally blank",
            "all rights reserved",
        ]
    ):
        return False
    if re.fullmatch(r"(?:l|c)\s+\d+/\d+\s+\d{1,2}\.\d{1,2}\.\d{4}", lowered):
        return False
    if profile == "scope_spec":
        return (
            "client to complete" in lowered
            or "required" in lowered
            or "testing must" in lowered
            or "evidence" in lowered
            or "approval" in lowered
            or "scope" in lowered
        )
    return bool(OBLIGATION_PATTERN.search(clean))


def extract_requirements(source: Source, document: DocumentRecord, text: str) -> list[Requirement]:
    rows: list[Requirement] = []
    seen: set[str] = set()
    created_at = now_iso()
    profile = source.parser_profile
    for citation, section, block in section_blocks(text, profile):
        for piece in split_sentences(block):
            piece = clean_line(piece)
            if not should_keep_requirement(profile, piece):
                continue
            key = f"{source.framework}|{citation}|{piece.lower()}"
            digest = hashlib.sha1(key.encode("utf-8")).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            obligation = normalized_obligation(piece)
            domain = infer_domain(profile, piece, section)
            response_type = infer_response_type(piece, domain)
            service_codes = infer_service_codes(piece, domain)
            rows.append(
                Requirement(
                    requirement_id=stable_id(source.framework, source.id, citation, piece),
                    source_id=source.id,
                    document_id=document.document_id,
                    framework=source.framework,
                    authority=source.authority,
                    jurisdiction=source.jurisdiction,
                    citation=citation,
                    section_title=section[:240],
                    domain=domain,
                    requirement_text=piece,
                    normalized_obligation=obligation,
                    evidence_expectation=evidence_expectation(profile, domain, piece),
                    vendor_question=vendor_question(source.framework, citation, obligation, domain),
                    response_type=response_type,
                    service_codes=service_codes,
                    source_hash=document.sha256 or "",
                    extraction_confidence=profile_confidence(profile, piece),
                    created_at=created_at,
                )
            )
    return rows


def document_text_and_record(source: Source) -> tuple[str, DocumentRecord] | None:
    path = source_file_path(source)
    if not path.is_file():
        return None
    suffix = path.suffix.lower()
    page_count: int | None = None
    if suffix == ".pdf":
        text, page_count = extract_pdf_text(path)
        doc_type = "pdf"
    elif suffix in {".html", ".htm"}:
        text = extract_html_text(path)
        doc_type = "html"
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
        doc_type = suffix.lstrip(".") or "text"
    file_hash = sha256_file(path)
    doc_id = stable_id("DOC", source.id, file_hash)
    record = DocumentRecord(
        document_id=doc_id,
        source_id=source.id,
        title=source.title,
        path=rel(path),
        doc_type=doc_type,
        sha256=file_hash,
        page_count=page_count,
        parser_profile=source.parser_profile,
        extracted_at=now_iso(),
    )
    return text, record


def insert_document(conn: sqlite3.Connection, document: DocumentRecord) -> None:
    conn.execute(
        """
        INSERT INTO documents (
            document_id, source_id, title, path, doc_type, sha256, page_count,
            parser_profile, extracted_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            document.document_id,
            document.source_id,
            document.title,
            document.path,
            document.doc_type,
            document.sha256,
            document.page_count,
            document.parser_profile,
            document.extracted_at,
        ),
    )


def insert_requirement(conn: sqlite3.Connection, requirement: Requirement) -> None:
    conn.execute(
        """
        INSERT OR IGNORE INTO requirements (
            requirement_id, source_id, document_id, framework, authority,
            jurisdiction, citation, section_title, domain, requirement_text,
            normalized_obligation, evidence_expectation, vendor_question,
            response_type, service_codes, source_hash, extraction_confidence,
            created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            requirement.requirement_id,
            requirement.source_id,
            requirement.document_id,
            requirement.framework,
            requirement.authority,
            requirement.jurisdiction,
            requirement.citation,
            requirement.section_title,
            requirement.domain,
            requirement.requirement_text,
            requirement.normalized_obligation,
            requirement.evidence_expectation,
            requirement.vendor_question,
            requirement.response_type,
            requirement.service_codes,
            requirement.source_hash,
            requirement.extraction_confidence,
            requirement.created_at,
        ),
    )
    for mapping_type, values in {
        "domain": [requirement.domain],
        "service_code": [value.strip() for value in requirement.service_codes.split(",")],
        "evidence": [requirement.evidence_expectation],
    }.items():
        for value in values:
            if not value:
                continue
            mapping_id = stable_id("MAP", requirement.requirement_id, mapping_type, value)
            conn.execute(
                """
                INSERT OR IGNORE INTO requirement_mappings (
                    mapping_id, requirement_id, mapping_type, mapping_value, rationale
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (
                    mapping_id,
                    requirement.requirement_id,
                    mapping_type,
                    value,
                    "Deterministic keyword mapping from regulatory requirement text.",
                ),
            )


def build_index(
    sources: list[Source],
    db_path: Path = DEFAULT_DB,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
) -> dict[str, Any]:
    ensure_dependencies("index")
    output_dir.mkdir(parents=True, exist_ok=True)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    if db_path.exists():
        db_path.unlink()
    download_results = load_download_results(output_dir)
    conn = sqlite3.connect(str(db_path))
    try:
        create_schema(conn)
        parsed_documents = 0
        extracted_requirements = 0
        for source in sources:
            insert_source(conn, source, download_results.get(source.id))
            extracted = document_text_and_record(source)
            if extracted is None:
                continue
            text, document = extracted
            insert_document(conn, document)
            parsed_documents += 1
            for requirement in extract_requirements(source, document, text):
                before = conn.total_changes
                insert_requirement(conn, requirement)
                if conn.total_changes > before:
                    extracted_requirements += 1
        conn.commit()
        summary = {
            "db_path": rel(db_path),
            "db_sha256": sha256_file(db_path),
            "sources": len(sources),
            "documents": parsed_documents,
            "requirements": extracted_requirements,
            "generated_at": now_iso(),
        }
        return summary
    finally:
        conn.close()


def rows(conn: sqlite3.Connection, query: str, params: tuple[Any, ...] = ()) -> list[sqlite3.Row]:
    conn.row_factory = sqlite3.Row
    return list(conn.execute(query, params))


def table_rows(conn: sqlite3.Connection, table: str) -> list[dict[str, Any]]:
    conn.row_factory = sqlite3.Row
    return [dict(row) for row in conn.execute(f"SELECT * FROM {table}")]


def autosize(ws: Any, max_width: int = 70) -> None:
    for column in ws.columns:
        letter = column[0].column_letter
        longest = 0
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            longest = max(longest, min(max_width, len(value) + 2))
        ws.column_dimensions[letter].width = max(10, min(max_width, longest))


def add_table(ws: Any, name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def style_sheet(ws: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    autosize(ws)


def write_rows(ws: Any, headers: list[str], data: Iterable[dict[str, Any]]) -> None:
    ws.append(headers)
    for item in data:
        ws.append([item.get(header) for header in headers])
    style_sheet(ws)


STATUS_OPTIONS = ["Not started", "In progress", "Submitted", "Complete", "Exception", "N/A"]
APPLICABILITY_OPTIONS = ["Vendor to confirm", "Applicable", "Not applicable", "Reference only", "Not provided", "Client owned"]
IMPLEMENTATION_OPTIONS = ["Not started", "Implemented", "Partially implemented", "Planned", "Compensating control", "Not applicable"]
EVIDENCE_STATUS_OPTIONS = ["Not provided", "Provided", "Referenced", "Pending", "Exception", "Not applicable"]
PRIORITY_OPTIONS = ["Critical", "High", "Medium", "Low", "Review"]
RESPONSE_REQUIRED_OPTIONS = ["Yes", "Review", "No"]


def header_map(ws: Any) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in ws[1] if cell.value is not None}


def set_tab_color(ws: Any, color: str) -> None:
    ws.sheet_properties.tabColor = color


def set_column_widths(ws: Any, widths: dict[str, float]) -> None:
    headers = header_map(ws)
    for name, width in widths.items():
        column = headers.get(name)
        if column:
            ws.column_dimensions[get_column_letter(column)].width = width


def add_dropdown(ws: Any, header: str, options: list[str]) -> None:
    if DataValidation is None or ws.max_row < 2:
        return
    column = header_map(ws).get(header)
    if not column:
        return
    letter = get_column_letter(column)
    formula = '"' + ",".join(options) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Select a value from the list."
    validation.errorTitle = "Invalid selection"
    validation.prompt = "Select a standard response value."
    validation.promptTitle = header.replace("_", " ").title()
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{max(ws.max_row, 5000)}")


def add_status_formatting(ws: Any, status_header: str = "completion_status") -> None:
    if FormulaRule is None or ws.max_row < 2:
        return
    headers = header_map(ws)
    status_col = headers.get(status_header)
    if not status_col:
        return
    status_letter = get_column_letter(status_col)
    body_range = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    fills = {
        "Complete": PatternFill("solid", fgColor="E2F0D9"),
        "Submitted": PatternFill("solid", fgColor="DDEBF7"),
        "In progress": PatternFill("solid", fgColor="FFF2CC"),
        "Exception": PatternFill("solid", fgColor="FCE4D6"),
        "Not started": PatternFill("solid", fgColor="F2F2F2"),
    }
    for value, fill in fills.items():
        ws.conditional_formatting.add(
            body_range,
            FormulaRule(formula=[f'${status_letter}2="{value}"'], fill=fill),
        )


def add_source_status_formatting(ws: Any, status_header: str = "status") -> None:
    if FormulaRule is None or ws.max_row < 2:
        return
    headers = header_map(ws)
    status_col = headers.get(status_header)
    if not status_col:
        return
    status_letter = get_column_letter(status_col)
    body_range = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    ok_values = ["downloaded", "cached", "available", "current", "updated"]
    for value in ok_values:
        ws.conditional_formatting.add(
            body_range,
            FormulaRule(formula=[f'${status_letter}2="{value}"'], fill=PatternFill("solid", fgColor="E2F0D9")),
        )
    ws.conditional_formatting.add(
        body_range,
        FormulaRule(formula=[f'${status_letter}2="fallback_available"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        body_range,
        FormulaRule(
            formula=[f'OR(${status_letter}2="failed",${status_letter}2="missing",${status_letter}2="cached_refresh_failed")'],
            fill=PatternFill("solid", fgColor="FCE4D6"),
        ),
    )


def add_quality_formatting(ws: Any, quality_header: str = "quality_flag") -> None:
    if FormulaRule is None or ws.max_row < 2:
        return
    headers = header_map(ws)
    quality_col = headers.get(quality_header)
    if not quality_col:
        return
    quality_letter = get_column_letter(quality_col)
    body_range = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.conditional_formatting.add(
        body_range,
        FormulaRule(formula=[f'${quality_letter}2="OK"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    ws.conditional_formatting.add(
        body_range,
        FormulaRule(formula=[f'${quality_letter}2<>"OK"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        body_range,
        FormulaRule(formula=[f'ISNUMBER(SEARCH("Source",${quality_letter}2))'], fill=PatternFill("solid", fgColor="FCE4D6")),
    )


def source_issue_type(source: dict[str, Any]) -> str:
    status = str(source.get("status") or "").lower()
    access = str(source.get("access") or "").lower()
    if status in {"failed", "missing"}:
        return "Source unavailable"
    if status == "fallback_available":
        return "Fallback source used"
    freshness = str(source.get("freshness_status") or "").lower()
    if freshness == "check_failed":
        return "Freshness check failed"
    source_page_status = str(source.get("source_page_status") or "").lower()
    if source_page_status == "check_failed":
        return "Source page freshness check failed"
    if not source.get("sha256"):
        return "Missing source hash"
    if "manual" in access or "restricted" in access:
        return "Manual/restricted source"
    return ""


def source_issue_action(source: dict[str, Any]) -> str:
    issue = source_issue_type(source)
    if issue == "Source unavailable":
        return "Re-check the official regulator page and add the current source if accessible."
    if issue == "Freshness check failed":
        return "Re-run refresh with network access and confirm the official source page manually if needed."
    if issue == "Source page freshness check failed":
        return "Confirm the regulator landing page manually and update the manifest if it points to a newer official document."
    if issue == "Fallback source used":
        return "Confirm the fallback archive is the intended official source for extraction."
    if issue == "Missing source hash":
        return "Refresh the local archive and rebuild the index so provenance can be hashed."
    if issue == "Manual/restricted source":
        return "Confirm access terms and attach the approved local source copy when available."
    return "No follow-up required."


def short_source_issue_action(issue: dict[str, Any]) -> str:
    issue_type = issue.get("issue_type")
    if issue_type == "Source unavailable":
        return "Re-check official source URL."
    if issue_type == "Freshness check failed":
        return "Re-run source refresh."
    if issue_type == "Source page freshness check failed":
        return "Confirm official landing page."
    if issue_type == "Fallback source used":
        return "Confirm fallback archive."
    if issue_type == "Missing source hash":
        return "Refresh local archive."
    if issue_type == "Manual/restricted source":
        return "Confirm access approval."
    return "No follow-up required."


def has_ocr_artifact(text: str) -> bool:
    suspicious = [
        "op tion",
        "infor m",
        "av ail",
        "theret o",
        "Offi cial",
        "Jour nal",
        "manag ement",
        "fi n",
        "im p",
        "propor tionality",
        "pr inciple",
        "chapt er",
        "relat ed",
        "fi nd",
        "im pact",
        "appropr iat",
        "ict -",
        "manag e",
        "repor te d",
        "secur ity",
        "integrit y",
        "strateg y",
    ]
    lowered = text.lower()
    return any(item.lower() in lowered for item in suspicious)


def priority_for_requirement(row: dict[str, Any], source: dict[str, Any] | None = None) -> str:
    domain = str(row.get("domain") or "")
    framework = str(row.get("framework") or "")
    status = str((source or {}).get("status") or "").lower()
    if status in {"failed", "missing"}:
        return "Review"
    if domain in {"Threat-led testing", "ICT risk management", "Third-party risk", "Operational resilience", "Incident response"}:
        return "High"
    if domain in {"Penetration testing", "Governance and approvals", "Scope and assets", "Evidence and reporting"}:
        return "Medium"
    if framework == "SCOPE_TEMPLATE":
        return "Medium"
    return "Low"


def evidence_artifact_type(row: dict[str, Any]) -> str:
    text = f"{row.get('domain', '')} {row.get('evidence_expectation', '')}".lower()
    if "threat intelligence" in text or "test plan" in text or "execution report" in text:
        return "Testing report package"
    if "provider register" in text or "contract" in text or "due-diligence" in text:
        return "Third-party control evidence"
    if "asset inventory" in text or "scope" in text:
        return "Inventory or scope register"
    if "incident" in text or "notification" in text:
        return "Incident response evidence"
    if "bc/dr" in text or "recovery" in text or "continuity" in text:
        return "Resilience plan and test record"
    if "completed vendor response" in text:
        return "Vendor questionnaire response"
    return "Policy, procedure, or control record"


def vendor_applicability(row: dict[str, Any]) -> str:
    if row.get("framework") == "SCOPE_TEMPLATE":
        return "Applicable"
    response_required = row.get("response_required") or response_required_for_requirement(row)
    if response_required == "No":
        return "Reference only"
    return "Vendor to confirm"


def response_required_for_requirement(row: dict[str, Any], source: dict[str, Any] | None = None) -> str:
    text = f"{row.get('section_title', '')} {row.get('normalized_obligation', '')} {row.get('requirement_text', '')}".lower()
    domain = str(row.get("domain") or "")
    framework = str(row.get("framework") or "")
    source_status = str((source or {}).get("status") or row.get("source_status") or "").lower()
    if source_status in {"failed", "missing"}:
        return "Review"
    if framework == "SCOPE_TEMPLATE":
        return "Yes"
    reference_only_patterns = [
        "definitions for the purposes",
        "for the purposes of this regulation",
        "the commission shall make that information publicly available",
        "member states may exclude",
        "where a member state makes use",
        "competent authorities shall decide",
        "lead overseers may",
        "the assessment referred to",
        "shall not apply to small",
        "this regulation applies to the following entities",
    ]
    if any(pattern in text for pattern in reference_only_patterns):
        return "No"
    if domain == "General regulatory obligation" and not re.search(
        r"\b(implement|maintain|establish|document|provide|conduct|perform|test|report|monitor|review|approve|evidence|retain|remediate|manage)\b",
        text,
    ):
        return "Review"
    if domain in {"Threat-led testing", "ICT risk management", "Third-party risk", "Operational resilience", "Incident response", "Penetration testing"}:
        return "Yes"
    if re.search(
        r"\b(shall|must|required|expected|provide|evidence|implement|maintain|conduct|perform|test|report|document|approve|review|monitor|remediation)\b",
        text,
    ):
        return "Yes"
    return "Review"


def vendor_disposition(row: dict[str, Any]) -> str:
    response_required = row.get("response_required") or response_required_for_requirement(row)
    if response_required == "Yes":
        return "Answer in workbook"
    if response_required == "Review":
        return "Review applicability"
    return "Reference only"


def vendor_workstream(row: dict[str, Any]) -> str:
    domain = str(row.get("domain") or "")
    if domain in {"Scope and assets", "Vendor intake and scope completion"}:
        return "Scope confirmation"
    if domain in {"Third-party risk", "Governance and approvals"}:
        return "Governance and third-party risk"
    if domain in {"Threat-led testing", "Penetration testing", "Tester qualification"}:
        return "Security testing"
    if domain in {"Incident response", "Operational resilience"}:
        return "Resilience and incident response"
    if domain == "ICT risk management":
        return "ICT control evidence"
    if domain == "Evidence and reporting":
        return "Evidence and reporting"
    return "Regulatory review"


def actionability_note(row: dict[str, Any]) -> str:
    response_required = row.get("response_required") or response_required_for_requirement(row)
    if response_required == "Yes":
        return "Vendor should provide a response and evidence reference."
    if response_required == "Review":
        return "Client/vendor should confirm applicability before completion."
    return "Kept for traceability; no vendor response expected by default."


def short_vendor_question(row: dict[str, Any]) -> str:
    obligation = str(row.get("normalized_obligation") or row.get("requirement_text") or "")
    if len(obligation) > 260:
        obligation = obligation[:257].rsplit(" ", 1)[0] + "..."
    return f"Provide implementation status, evidence, and any gaps for: {obligation}"


def quality_assessment(row: dict[str, Any], source: dict[str, Any] | None = None) -> tuple[str, str]:
    notes: list[str] = []
    source = source or {}
    issue = source_issue_type(source) if source else ""
    if issue:
        notes.append(issue)
    if not row.get("source_hash"):
        notes.append("Missing requirement source hash")
    confidence = float(row.get("extraction_confidence") or 0)
    if confidence < 0.75:
        notes.append("Lower parser confidence")
    text = str(row.get("requirement_text") or "")
    if len(text) < 60:
        notes.append("Short requirement text")
    if has_ocr_artifact(text):
        notes.append("Potential OCR artifact")
    if str(row.get("citation") or "").lower() == "general":
        notes.append("General citation")
    if row.get("domain") == "General regulatory obligation":
        notes.append("Broad domain classification")
    if not notes:
        return "OK", "No automated quality issues detected."
    if any("Source" in note or "hash" in note for note in notes):
        return "Source review", "; ".join(dict.fromkeys(notes))
    if any("OCR" in note or "Short" in note or "General citation" in note for note in notes):
        return "Content review", "; ".join(dict.fromkeys(notes))
    return "Classification review", "; ".join(dict.fromkeys(notes))


def enriched_requirement_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    source_lookup = {item["source_id"]: item for item in table_rows(conn, "sources")}
    out: list[dict[str, Any]] = []
    for item in table_rows(conn, "requirements"):
        source = source_lookup.get(item["source_id"], {})
        quality_flag, quality_note = quality_assessment(item, source)
        row = dict(item)
        response_required = response_required_for_requirement(item, source)
        row.update(
            {
                "source_title": source.get("title", ""),
                "source_status": source.get("status", ""),
                "priority": priority_for_requirement(item, source),
                "quality_flag": quality_flag,
                "quality_note": quality_note,
                "evidence_artifact_type": evidence_artifact_type(item),
                "response_required": response_required,
                "vendor_disposition": vendor_disposition({"response_required": response_required}),
                "vendor_workstream": vendor_workstream(item),
                "vendor_short_question": short_vendor_question(item),
                "actionability_note": actionability_note({"response_required": response_required}),
                "vendor_applicability": vendor_applicability({"framework": item["framework"], "response_required": response_required}),
            }
        )
        out.append(row)
    return out


def source_issue_rows(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for source in table_rows(conn, "sources"):
        issue = source_issue_type(source)
        if not issue:
            continue
        issues.append(
            {
                "source_id": source["source_id"],
                "framework": source["framework"],
                "status": source["status"],
                "issue_type": issue,
                "recommended_action": source_issue_action(source),
                "title": source["title"],
                "local_path": source["local_path"],
                "source_url": source["source_url"],
                "source_page_status": source.get("source_page_status", ""),
                "freshness_status": source.get("freshness_status", ""),
                "notes": source["notes"],
            }
        )
    return issues


def count_values(items: Iterable[dict[str, Any]], key: str) -> dict[str, int]:
    counts: dict[str, int] = {}
    for item in items:
        value = str(item.get(key) or "Unknown")
        counts[value] = counts.get(value, 0) + 1
    return dict(sorted(counts.items(), key=lambda pair: (-pair[1], pair[0])))


def style_dashboard_cell(cell: Any, fill: str, font_color: str = "0B2545", bold: bool = False, size: int = 10) -> None:
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=font_color, bold=bold, size=size)
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def write_dashboard_table(ws: Any, start_row: int, start_col: int, title: str, headers: list[str], data: list[list[Any]]) -> int:
    ws.cell(start_row, start_col, title)
    ws.merge_cells(
        start_row=start_row,
        start_column=start_col,
        end_row=start_row,
        end_column=start_col + len(headers) - 1,
    )
    style_dashboard_cell(ws.cell(start_row, start_col), "DDEBF7", bold=True, size=11)
    header_row = start_row + 1
    for index, header in enumerate(headers, start=start_col):
        cell = ws.cell(header_row, index, header)
        style_dashboard_cell(cell, "1F4E78", "FFFFFF", True)
    for row_index, values in enumerate(data, start=header_row + 1):
        for col_index, value in enumerate(values, start=start_col):
            cell = ws.cell(row_index, col_index, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=Side(style="thin", color="D9E2F3"))
    return header_row + len(data)


def build_dashboard_sheet(
    wb: Any,
    conn: sqlite3.Connection,
    enriched_rows: list[dict[str, Any]],
    issues: list[dict[str, Any]],
    db_path: Path,
    db_hash: str,
) -> Any:
    ws = wb.active
    ws.title = "Dashboard"
    set_tab_color(ws, "1F4E78")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "Regulatory Vendor Requirements Pack"
    ws["A1"].fill = PatternFill("solid", fgColor="0B2545")
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:H2")
    ws["A2"] = "Traceable vendor questionnaire generated from index.db, official source metadata, and supplied scope documents."
    ws["A2"].font = Font(color="44546A", italic=True)

    metadata = [
        ("Generated at", now_iso()),
        ("Database path", rel(db_path)),
        ("Database SHA-256", f"{db_hash[:16]}...{db_hash[-8:]}" if db_hash else ""),
        ("Known source issues", len(issues)),
    ]
    for row_offset, (label, value) in enumerate(metadata, start=4):
        ws.cell(row_offset, 1, label)
        ws.cell(row_offset, 2, value)
        style_dashboard_cell(ws.cell(row_offset, 1), "E8EEF5", bold=True)
        ws.cell(row_offset, 2).alignment = Alignment(wrap_text=True, vertical="top")

    requirement_count = len(enriched_rows)
    source_count = conn.execute("SELECT COUNT(*) FROM sources").fetchone()[0]
    high_priority = sum(1 for row in enriched_rows if row["priority"] in {"Critical", "High"})
    response_required = sum(1 for row in enriched_rows if row["response_required"] == "Yes")
    review_rows = sum(1 for row in enriched_rows if row["quality_flag"] != "OK")
    kpis = [
        ("Requirements", requirement_count, "All clause-level rows in index.db"),
        ("Vendor worklist", response_required, "Rows expecting a vendor answer"),
        ("High priority", high_priority, "Rows vendors should answer early"),
        ("Review flags", review_rows, "Rows needing source/content review"),
        ("Sources", source_count, "Official and supplied source records"),
    ]
    ws.merge_cells("D4:H4")
    ws["D4"] = "Vendor Pack Metrics"
    style_dashboard_cell(ws["D4"], "1F4E78", "FFFFFF", True, 11)
    for index, header in enumerate(["Metric", "Count", "Why it matters"], start=4):
        cell = ws.cell(5, index, header)
        style_dashboard_cell(cell, "DDEBF7", "0B2545", True)
    for row_index, (label, value, note) in enumerate(kpis, start=6):
        ws.cell(row_index, 4, label)
        ws.cell(row_index, 5, value)
        ws.cell(row_index, 6, note)
        ws.merge_cells(start_row=row_index, start_column=6, end_row=row_index, end_column=8)
        style_dashboard_cell(ws.cell(row_index, 4), "F8FBFD", "0B2545", True)
        style_dashboard_cell(ws.cell(row_index, 5), "F8FBFD", "0B2545", True)
        ws.cell(row_index, 6).alignment = Alignment(wrap_text=True, vertical="top")

    framework_data = [[key, value] for key, value in count_values(enriched_rows, "framework").items()]
    write_dashboard_table(ws, 12, 1, "Requirements by Framework", ["Framework", "Rows"], framework_data)

    source_status_rows = rows(conn, "SELECT status, COUNT(*) AS count FROM sources GROUP BY status ORDER BY count DESC, status")
    write_dashboard_table(ws, 12, 4, "Source Status", ["Status", "Sources"], [[r["status"], r["count"]] for r in source_status_rows])

    qa_data = [[key, value] for key, value in count_values(enriched_rows, "quality_flag").items()]
    write_dashboard_table(ws, 12, 7, "Quality Flags", ["Flag", "Rows"], qa_data)

    actionability_data = [[key, value] for key, value in count_values(enriched_rows, "response_required").items()]
    write_dashboard_table(ws, 23, 1, "Vendor Response Triage", ["Response Required", "Rows"], actionability_data)

    workstream_data = [[key, value] for key, value in list(count_values(enriched_rows, "vendor_workstream").items())[:8]]
    write_dashboard_table(ws, 23, 4, "Top Vendor Workstreams", ["Workstream", "Rows"], workstream_data)

    issue_preview = [[key, value] for key, value in count_values(issues, "issue_type").items()] or [["No known source issues", 0]]
    write_dashboard_table(
        ws,
        34,
        1,
        "Known Source Issues",
        ["Issue Type", "Sources"],
        issue_preview,
    )

    widths = {"A": 24, "B": 18, "C": 18, "D": 24, "E": 14, "F": 28, "G": 18, "H": 18}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A12"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def vendor_completion_headers(include_framework: bool = True) -> list[str]:
    headers = [
        "requirement_id",
        "citation",
        "domain",
        "priority",
        "response_required",
        "applicability",
        "vendor_workstream",
        "vendor_short_question",
        "evidence_artifact_type",
        "implementation_status",
        "evidence_status",
        "vendor_response",
        "evidence_reference",
        "control_owner",
        "gap_or_exception_notes",
        "remediation_owner",
        "target_date",
        "completion_status",
    ]
    if include_framework:
        headers.insert(1, "framework")
    return headers


def vendor_completion_row(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "requirement_id": item["requirement_id"],
        "framework": item["framework"],
        "citation": item["citation"],
        "domain": item["domain"],
        "priority": item["priority"],
        "response_required": item["response_required"],
        "applicability": item["vendor_applicability"],
        "vendor_workstream": item["vendor_workstream"],
        "vendor_short_question": item["vendor_short_question"],
        "evidence_artifact_type": item["evidence_artifact_type"],
        "implementation_status": "Not started",
        "evidence_status": "Not provided",
        "vendor_response": "",
        "evidence_reference": "",
        "control_owner": "",
        "gap_or_exception_notes": "",
        "remediation_owner": "",
        "target_date": "",
        "completion_status": "Not started",
    }


def style_vendor_completion_sheet(ws: Any, table_name: str) -> None:
    add_table(ws, table_name)
    add_dropdown(ws, "priority", PRIORITY_OPTIONS)
    add_dropdown(ws, "response_required", RESPONSE_REQUIRED_OPTIONS)
    add_dropdown(ws, "applicability", APPLICABILITY_OPTIONS)
    add_dropdown(ws, "implementation_status", IMPLEMENTATION_OPTIONS)
    add_dropdown(ws, "evidence_status", EVIDENCE_STATUS_OPTIONS)
    add_dropdown(ws, "completion_status", STATUS_OPTIONS)
    add_status_formatting(ws, "completion_status")
    set_column_widths(
        ws,
        {
            "requirement_id": 20,
            "framework": 18,
            "citation": 16,
            "domain": 24,
            "priority": 13,
            "response_required": 18,
            "applicability": 19,
            "vendor_workstream": 26,
            "vendor_short_question": 62,
            "evidence_artifact_type": 30,
            "implementation_status": 22,
            "evidence_status": 18,
            "vendor_response": 44,
            "evidence_reference": 34,
            "gap_or_exception_notes": 40,
        },
    )
    for row_index in range(2, ws.max_row + 1):
        ws.row_dimensions[row_index].height = 48


def add_vendor_completion_sheet(
    wb: Any,
    sheet_name: str,
    table_name: str,
    rows_for_sheet: list[dict[str, Any]],
    include_framework: bool = True,
) -> Any:
    ws = wb.create_sheet(sheet_name)
    set_tab_color(ws, "ED7D31")
    headers = vendor_completion_headers(include_framework=include_framework)
    write_rows(ws, headers, [vendor_completion_row(item) for item in rows_for_sheet])
    style_vendor_completion_sheet(ws, table_name)
    return ws


def evidence_checklist_rows(enriched_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str], dict[str, Any]] = {}
    for item in enriched_rows:
        if item["response_required"] == "No":
            continue
        key = (item["framework"], item["domain"], item["evidence_artifact_type"])
        record = grouped.setdefault(
            key,
            {
                "framework": item["framework"],
                "domain": item["domain"],
                "vendor_workstream": item["vendor_workstream"],
                "evidence_artifact_type": item["evidence_artifact_type"],
                "requirement_count": 0,
                "highest_priority": item["priority"],
                "evidence_expectation": item["evidence_expectation"],
                "service_codes": set(),
                "evidence_owner": "",
                "evidence_status": "Not provided",
                "evidence_reference": "",
                "notes": "",
            },
        )
        record["requirement_count"] += 1
        if PRIORITY_OPTIONS.index(item["priority"]) < PRIORITY_OPTIONS.index(record["highest_priority"]):
            record["highest_priority"] = item["priority"]
        for code in str(item["service_codes"]).split(","):
            cleaned = code.strip()
            if cleaned:
                record["service_codes"].add(cleaned)
    out: list[dict[str, Any]] = []
    for record in grouped.values():
        row = dict(record)
        row["service_codes"] = ", ".join(sorted(row["service_codes"]))
        out.append(row)
    return sorted(out, key=lambda item: (item["framework"], item["vendor_workstream"], item["domain"], item["evidence_artifact_type"]))


def export_xlsx(conn: sqlite3.Connection, output_dir: Path, db_path: Path) -> Path:
    if Workbook is None:
        raise RuntimeError("openpyxl is required to export XLSX")
    output_dir.mkdir(parents=True, exist_ok=True)
    dated = output_dir / f"regulatory_requirements_reference_{today_slug()}.xlsx"
    latest = output_dir / "regulatory_requirements_reference_latest.xlsx"
    requirement_count = conn.execute("SELECT COUNT(*) FROM requirements").fetchone()[0]
    source_count = conn.execute("SELECT COUNT(*) FROM sources").fetchone()[0]
    db_hash = sha256_file(db_path) if db_path.exists() else ""
    enriched_rows = enriched_requirement_rows(conn)
    issues = source_issue_rows(conn)

    wb = Workbook()
    build_dashboard_sheet(wb, conn, enriched_rows, issues, db_path, db_hash)

    readme = wb.create_sheet("Read Me")
    set_tab_color(readme, "5B9BD5")
    readme.append(["Field", "Value"])
    readme_rows = [
        ("Generated at", now_iso()),
        ("Database path", rel(db_path)),
        ("Database SHA-256", db_hash),
        ("Requirement rows", requirement_count),
        ("Source rows", source_count),
        ("Known source issues", len(issues)),
        ("Usage", "Use Dashboard to triage, Vendor Worklist for active completion, framework vendor tabs for assignment, QA Flags for review, and Requirements for audit traceability."),
    ]
    for row in readme_rows:
        readme.append(list(row))
    style_sheet(readme)

    source_headers = [
        "source_id",
        "title",
        "authority",
        "framework",
        "jurisdiction",
        "source_url",
        "download_url",
        "local_path",
        "source_kind",
        "access",
        "status",
        "sha256",
        "latest_url",
        "latest_sha256",
        "source_page_url",
        "source_page_sha256",
        "source_page_status",
        "source_page_checked_at",
        "freshness_status",
        "freshness_checked_at",
        "notes",
    ]
    ws = wb.create_sheet("Sources")
    set_tab_color(ws, "70AD47")
    write_rows(ws, source_headers, table_rows(conn, "sources"))
    add_table(ws, "SourcesTable")
    add_source_status_formatting(ws, "status")
    set_column_widths(
        ws,
        {
            "title": 42,
            "source_url": 48,
            "download_url": 48,
            "local_path": 44,
            "notes": 52,
            "sha256": 28,
            "latest_url": 48,
            "latest_sha256": 28,
            "source_page_url": 48,
            "source_page_sha256": 28,
            "source_page_status": 22,
            "source_page_checked_at": 24,
            "freshness_status": 20,
            "freshness_checked_at": 24,
        },
    )

    requirement_headers = [
        "requirement_id",
        "framework",
        "authority",
        "jurisdiction",
        "citation",
        "section_title",
        "domain",
        "requirement_text",
        "normalized_obligation",
        "evidence_expectation",
        "vendor_question",
        "response_type",
        "service_codes",
        "source_id",
        "source_title",
        "source_status",
        "source_hash",
        "extraction_confidence",
        "priority",
        "response_required",
        "vendor_disposition",
        "vendor_applicability",
        "vendor_workstream",
        "evidence_artifact_type",
        "actionability_note",
        "quality_flag",
        "quality_note",
    ]
    ws = wb.create_sheet("Requirements")
    set_tab_color(ws, "4472C4")
    write_rows(ws, requirement_headers, enriched_rows)
    add_table(ws, "RequirementsTable")
    add_dropdown(ws, "priority", PRIORITY_OPTIONS)
    add_dropdown(ws, "vendor_applicability", APPLICABILITY_OPTIONS)
    add_quality_formatting(ws, "quality_flag")
    set_column_widths(
        ws,
        {
            "requirement_id": 20,
            "authority": 28,
            "citation": 16,
            "section_title": 28,
            "domain": 24,
            "requirement_text": 56,
            "normalized_obligation": 48,
            "evidence_expectation": 45,
            "vendor_question": 58,
            "service_codes": 24,
            "source_title": 36,
            "source_hash": 28,
            "quality_note": 48,
        },
    )

    vendor_headers = [
        "requirement_id",
        "framework",
        "citation",
        "domain",
        "priority",
        "response_required",
        "vendor_disposition",
        "applicability",
        "vendor_workstream",
        "vendor_question",
        "vendor_short_question",
        "response_type",
        "evidence_artifact_type",
        "implementation_status",
        "evidence_status",
        "vendor_response",
        "evidence_reference",
        "control_owner",
        "gap_or_exception_notes",
        "remediation_owner",
        "target_date",
        "reviewer_notes",
        "completion_status",
    ]
    vendor_rows = []
    for item in enriched_rows:
        vendor_rows.append(
            {
                "requirement_id": item["requirement_id"],
                "framework": item["framework"],
                "citation": item["citation"],
                "domain": item["domain"],
                "priority": item["priority"],
                "response_required": item["response_required"],
                "vendor_disposition": item["vendor_disposition"],
                "applicability": item["vendor_applicability"],
                "vendor_workstream": item["vendor_workstream"],
                "vendor_question": item["vendor_question"],
                "vendor_short_question": item["vendor_short_question"],
                "response_type": item["response_type"],
                "evidence_artifact_type": item["evidence_artifact_type"],
                "implementation_status": "Not started",
                "evidence_status": "Not provided",
                "vendor_response": "",
                "evidence_reference": "",
                "control_owner": "",
                "gap_or_exception_notes": "",
                "remediation_owner": "",
                "target_date": "",
                "reviewer_notes": "",
                "completion_status": "Not started",
            }
        )
    ws = wb.create_sheet("Vendor Questions")
    set_tab_color(ws, "ED7D31")
    write_rows(ws, vendor_headers, vendor_rows)
    add_table(ws, "VendorQuestionsTable")
    add_dropdown(ws, "priority", PRIORITY_OPTIONS)
    add_dropdown(ws, "response_required", RESPONSE_REQUIRED_OPTIONS)
    add_dropdown(ws, "applicability", APPLICABILITY_OPTIONS)
    add_dropdown(ws, "response_type", ["narrative", "upload_and_narrative", "table", "date_or_frequency", "yes_no"])
    add_dropdown(ws, "implementation_status", IMPLEMENTATION_OPTIONS)
    add_dropdown(ws, "evidence_status", EVIDENCE_STATUS_OPTIONS)
    add_dropdown(ws, "completion_status", STATUS_OPTIONS)
    add_status_formatting(ws, "completion_status")
    set_column_widths(
        ws,
        {
            "requirement_id": 20,
            "citation": 16,
            "domain": 24,
            "priority": 13,
            "response_required": 18,
            "vendor_disposition": 22,
            "applicability": 20,
            "vendor_workstream": 26,
            "vendor_question": 64,
            "vendor_short_question": 62,
            "response_type": 22,
            "evidence_artifact_type": 30,
            "implementation_status": 22,
            "evidence_status": 18,
            "vendor_response": 48,
            "evidence_reference": 36,
            "gap_or_exception_notes": 42,
            "reviewer_notes": 36,
        },
    )
    for row_index in range(2, ws.max_row + 1):
        ws.row_dimensions[row_index].height = 54

    worklist_rows = [item for item in enriched_rows if item["response_required"] != "No"]
    add_vendor_completion_sheet(wb, "Vendor Worklist", "VendorWorklistTable", worklist_rows)

    vendor_sheet_groups = [
        ("Vendor - DORA", "VendorDORATable", {"DORA", "DORA_TLPT"}),
        ("Vendor - US Bank", "VendorUSBankTable", {"FRB_SUPERVISION", "FFIEC"}),
        ("Vendor - HKMA CBEST", "VendorHKMACBESTTable", {"HKMA_ICAST", "PRA_CBEST"}),
        ("Vendor - Scope", "VendorScopeTable", {"SCOPE_TEMPLATE"}),
    ]
    for sheet_name, table_name, frameworks in vendor_sheet_groups:
        add_vendor_completion_sheet(
            wb,
            sheet_name,
            table_name,
            [item for item in worklist_rows if item["framework"] in frameworks],
            include_framework=True,
        )

    evidence_query = """
        SELECT framework, domain, evidence_expectation, COUNT(*) AS requirement_count,
               GROUP_CONCAT(DISTINCT service_codes) AS service_codes
        FROM requirements
        GROUP BY framework, domain, evidence_expectation
        ORDER BY framework, domain
    """
    ws = wb.create_sheet("Evidence Matrix")
    set_tab_color(ws, "A5A5A5")
    write_rows(ws, ["framework", "domain", "evidence_expectation", "requirement_count", "service_codes"], [dict(row) for row in rows(conn, evidence_query)])
    add_table(ws, "EvidenceMatrixTable")
    set_column_widths(ws, {"domain": 24, "evidence_expectation": 58, "service_codes": 42})

    evidence_checklist_headers = [
        "framework",
        "domain",
        "vendor_workstream",
        "evidence_artifact_type",
        "requirement_count",
        "highest_priority",
        "evidence_expectation",
        "service_codes",
        "evidence_owner",
        "evidence_status",
        "evidence_reference",
        "notes",
    ]
    ws = wb.create_sheet("Evidence Checklist")
    set_tab_color(ws, "70AD47")
    write_rows(ws, evidence_checklist_headers, evidence_checklist_rows(enriched_rows))
    add_table(ws, "EvidenceChecklistTable")
    add_dropdown(ws, "highest_priority", PRIORITY_OPTIONS)
    add_dropdown(ws, "evidence_status", EVIDENCE_STATUS_OPTIONS)
    set_column_widths(
        ws,
        {
            "domain": 24,
            "vendor_workstream": 28,
            "evidence_artifact_type": 30,
            "evidence_expectation": 58,
            "service_codes": 34,
            "evidence_owner": 24,
            "evidence_reference": 38,
            "notes": 38,
        },
    )

    service_query = """
        SELECT mapping_value AS service_code, COUNT(*) AS requirement_count,
               GROUP_CONCAT(DISTINCT r.framework) AS frameworks,
               GROUP_CONCAT(DISTINCT r.domain) AS domains
        FROM requirement_mappings m
        JOIN requirements r ON r.requirement_id = m.requirement_id
        WHERE mapping_type = 'service_code'
        GROUP BY mapping_value
        ORDER BY mapping_value
    """
    ws = wb.create_sheet("Service Mapping")
    set_tab_color(ws, "A5A5A5")
    write_rows(ws, ["service_code", "requirement_count", "frameworks", "domains"], [dict(row) for row in rows(conn, service_query)])
    add_table(ws, "ServiceMappingTable")
    set_column_widths(ws, {"frameworks": 42, "domains": 54})

    qa_headers = [
        "requirement_id",
        "framework",
        "citation",
        "domain",
        "priority",
        "quality_flag",
        "quality_note",
        "source_id",
        "source_status",
        "requirement_text",
    ]
    qa_rows = [row for row in enriched_rows if row["quality_flag"] != "OK"]
    ws = wb.create_sheet("QA Flags")
    set_tab_color(ws, "FFC000")
    write_rows(ws, qa_headers, qa_rows)
    add_table(ws, "QAFlagsTable")
    add_dropdown(ws, "priority", PRIORITY_OPTIONS)
    add_quality_formatting(ws, "quality_flag")
    set_column_widths(ws, {"quality_note": 52, "requirement_text": 62})

    source_issue_headers = [
        "source_id",
        "framework",
        "status",
        "issue_type",
        "recommended_action",
        "title",
        "local_path",
        "source_url",
        "source_page_status",
        "freshness_status",
        "notes",
    ]
    ws = wb.create_sheet("Source Issues")
    set_tab_color(ws, "C00000")
    write_rows(ws, source_issue_headers, issues)
    add_table(ws, "SourceIssuesTable")
    set_column_widths(
        ws,
        {
            "source_id": 34,
            "recommended_action": 58,
            "title": 42,
            "local_path": 46,
            "source_url": 52,
            "source_page_status": 22,
            "freshness_status": 20,
            "notes": 50,
        },
    )

    dictionary_rows = [
        {"field": "requirement_id", "description": "Stable deterministic requirement identifier."},
        {"field": "citation", "description": "Article, section, phase, or generated location marker."},
        {"field": "source_hash", "description": "SHA-256 hash of the local source document at extraction time."},
        {"field": "latest_sha256", "description": "SHA-256 hash observed during the most recent online refresh of the archived source."},
        {"field": "source_page_sha256", "description": "SHA-256 hash of the regulator landing page checked during refresh, when different from the archive URL."},
        {"field": "freshness_status", "description": "Downloader freshness result for the archived source: current, updated, downloaded, not_checked, check_failed, offline_skipped, or workspace_supplied."},
        {"field": "extraction_confidence", "description": "Deterministic parser confidence score based on source profile and obligation language."},
        {"field": "priority", "description": "Export-derived vendor triage priority based on framework/domain/source status."},
        {"field": "response_required", "description": "Export-derived vendor actionability: Yes, Review, or No."},
        {"field": "vendor_disposition", "description": "Plain-English handling instruction for each row."},
        {"field": "vendor_workstream", "description": "Operational grouping used to route vendor completion work."},
        {"field": "vendor_short_question", "description": "Shorter completion prompt used by the vendor worklist tabs."},
        {"field": "vendor_applicability", "description": "Default applicability prompt for the vendor to confirm or override."},
        {"field": "quality_flag", "description": "Automated review signal for source, OCR, content, or classification checks."},
        {"field": "implementation_status", "description": "Vendor-selected current implementation state."},
        {"field": "evidence_status", "description": "Vendor-selected evidence availability state."},
        {"field": "vendor_response", "description": "Blank vendor completion field."},
        {"field": "evidence_reference", "description": "Vendor-provided attachment name, URL, GRC reference, or evidence pointer."},
    ]
    ws = wb.create_sheet("Data Dictionary")
    set_tab_color(ws, "5B9BD5")
    write_rows(ws, ["field", "description"], dictionary_rows)

    ws = wb.create_sheet("Change Log")
    set_tab_color(ws, "5B9BD5")
    write_rows(
        ws,
        ["generated_at", "change", "requirement_count", "source_count", "database_hash"],
        [
            {
                "generated_at": now_iso(),
                "change": "Generated vendor-ready regulatory requirements workbook from index.db with dashboard, QA, source issue, and vendor completion views.",
                "requirement_count": requirement_count,
                "source_count": source_count,
                "database_hash": db_hash,
            }
        ],
    )

    wb.save(dated)
    shutil.copyfile(dated, latest)
    return dated


QUESTIONNAIRE_DOMAIN_RESPONSE_PATHS: dict[str, dict[str, str]] = {
    "General regulatory obligation": {
        "code": "REG-01",
        "section": "Regulatory applicability and residual obligations",
        "question": "Confirm whether the proposed test scope, evidence package, or exception register addresses this regulator/domain grouping.",
    },
    "Threat-led testing": {
        "code": "TEST-02",
        "section": "Threat-led and advanced testing",
        "question": "Describe the threat-led testing model, control-team governance, scenarios, tester independence, and regulator/client approvals.",
    },
    "Third-party risk": {
        "code": "TPR-01",
        "section": "Third-party and subcontractor coverage",
        "question": "Identify all third-party, subcontractor, platform, and outsourced-service dependencies covered by the test and evidence package.",
    },
    "ICT risk management": {
        "code": "ICT-01",
        "section": "ICT control evidence",
        "question": "Provide the ICT governance, control ownership, policy, procedure, and assurance evidence supporting the proposed test scope.",
    },
    "Scope and assets": {
        "code": "SCOPE-01",
        "section": "Scope boundary and assets",
        "question": "Define in-scope assets, systems, applications, APIs, networks, cloud services, data classes, and explicit exclusions.",
    },
    "Evidence and reporting": {
        "code": "EVID-01",
        "section": "Evidence, reporting, and retention",
        "question": "Describe the deliverables, evidence references, retention location, reporting timeline, and finding/remediation tracking approach.",
    },
    "Operational resilience": {
        "code": "RES-01",
        "section": "Operational resilience and recovery",
        "question": "Explain how the test scope supports critical services, recovery expectations, continuity planning, and resilience evidence.",
    },
    "Incident response": {
        "code": "RES-02",
        "section": "Incident response and communications",
        "question": "Describe monitoring, escalation, incident communication, notification, containment, and lessons-learned evidence for the test.",
    },
    "Penetration testing": {
        "code": "TEST-01",
        "section": "Penetration testing execution",
        "question": "Define the penetration testing methodology, test types, access model, rules of engagement, safety controls, and reporting outputs.",
    },
    "Governance and approvals": {
        "code": "GOV-01",
        "section": "Governance, approvals, and accountability",
        "question": "Identify approving parties, control owners, legal/regulatory constraints, change approvals, and final acceptance responsibilities.",
    },
    "Tester qualification": {
        "code": "TEST-03",
        "section": "Tester qualification and independence",
        "question": "Provide tester qualifications, independence confirmations, certifications, subcontractor details, and conflict checks.",
    },
    "Vendor intake and scope completion": {
        "code": "INTAKE-01",
        "section": "Vendor intake and scope completion",
        "question": "Complete the vendor intake fields, reviewer-ready scope decisions, supporting attachments, and open completion items.",
    },
}

QUESTIONNAIRE_CORE_ROWS: list[dict[str, str]] = [
    {
        "code": "SCOPE-01",
        "coverage": "Scope and assets; vendor intake and scope completion",
        "prompt": "Define the exact systems, applications, APIs, networks, cloud environments, data classes, locations, and exclusions for this test.",
    },
    {
        "code": "TEST-01",
        "coverage": "Penetration testing and standard security testing",
        "prompt": "Describe the methodology, test types, access model, test windows, safety limits, rules of engagement, and deliverables.",
    },
    {
        "code": "TEST-02",
        "coverage": "Threat-led testing, TLPT, CBEST, iCAST, and advanced red-team expectations",
        "prompt": "Describe threat-intelligence inputs, scenario governance, control-team approvals, detection/response scope, and regulator/client oversight.",
    },
    {
        "code": "TEST-03",
        "coverage": "Tester qualification, independence, certification, and subcontractor suitability",
        "prompt": "Provide tester roles, qualifications, independence statements, subcontractor details, and any certification or conflict evidence.",
    },
    {
        "code": "TPR-01",
        "coverage": "Third-party, outsourced service, subcontractor, ICT provider, and supply-chain obligations",
        "prompt": "List third parties and outsourced providers in scope, contractual constraints, approval dependencies, and monitoring evidence.",
    },
    {
        "code": "ICT-01",
        "coverage": "ICT risk management, policies, procedures, control ownership, and assurance evidence",
        "prompt": "Identify control owners, policies, procedures, assurance records, ICT risk decisions, and evidence that supports the proposed test.",
    },
    {
        "code": "RES-01",
        "coverage": "Operational resilience, continuity, recovery, and critical service evidence",
        "prompt": "Describe critical services, recovery assumptions, continuity controls, resilience tests, and evidence affected by this assessment.",
    },
    {
        "code": "RES-02",
        "coverage": "Incident response, communications, notification, detection, and lessons learned",
        "prompt": "Describe monitoring, escalation, notification, communication, containment, and post-test lessons-learned evidence.",
    },
    {
        "code": "GOV-01",
        "coverage": "Governance, management approvals, regulator/client approval gates, and accountability",
        "prompt": "Identify approvers, control owners, legal constraints, regulator/client checkpoints, and final acceptance responsibilities.",
    },
    {
        "code": "EVID-01",
        "coverage": "Evidence, reporting, retention, remediation, and audit traceability",
        "prompt": "List the required deliverables, evidence repository, retention path, report sections, finding tracker, and remediation workflow.",
    },
    {
        "code": "REG-01",
        "coverage": "General regulatory obligations and reference/review rows",
        "prompt": "Confirm applicability of residual regulatory rows and document any non-applicable items, assumptions, exceptions, or review dependencies.",
    },
    {
        "code": "INTAKE-01",
        "coverage": "Vendor intake, reviewer workflow, and scope completion requirements",
        "prompt": "Complete the vendor profile, engagement metadata, reviewer routing, required attachments, and outstanding completion items.",
    },
]


def questionnaire_response_path(row: dict[str, Any]) -> dict[str, str] | None:
    return QUESTIONNAIRE_DOMAIN_RESPONSE_PATHS.get(str(row.get("domain") or ""))


def compact_join(values: Iterable[Any], limit: int = 5) -> str:
    cleaned = [str(value) for value in dict.fromkeys(value for value in values if value)]
    if len(cleaned) <= limit:
        return ", ".join(cleaned)
    return ", ".join(cleaned[:limit]) + f", +{len(cleaned) - limit} more"


def compact_requirement_ids(items: list[dict[str, Any]], limit: int = 4) -> str:
    ids = sorted(str(item["requirement_id"]) for item in items if item.get("requirement_id"))
    if len(ids) <= limit:
        return ", ".join(ids)
    return ", ".join(ids[:limit]) + f", +{len(ids) - limit} more"


def questionnaire_coverage_rows(enriched_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    grouped: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    unmapped: list[dict[str, Any]] = []
    for item in enriched_rows:
        path = questionnaire_response_path(item)
        if path is None:
            unmapped.append(item)
            continue
        key = (item["framework"], item["domain"], item["evidence_artifact_type"], path["code"])
        record = grouped.setdefault(
            key,
            {
                "framework": item["framework"],
                "domain": item["domain"],
                "evidence_artifact_type": item["evidence_artifact_type"],
                "response_path": path["code"],
                "response_section": path["section"],
                "vendor_question": path["question"],
                "requirement_count": 0,
                "response_required": set(),
                "priorities": set(),
                "items": [],
            },
        )
        record["requirement_count"] += 1
        record["response_required"].add(item["response_required"])
        record["priorities"].add(item["priority"])
        record["items"].append(item)
    out: list[dict[str, Any]] = []
    for record in grouped.values():
        row = dict(record)
        items = row.pop("items")
        row["response_required"] = compact_join(sorted(row["response_required"]))
        row["priorities"] = compact_join(sorted(row["priorities"], key=lambda value: PRIORITY_OPTIONS.index(value) if value in PRIORITY_OPTIONS else 99))
        row["requirement_ids"] = compact_requirement_ids(items)
        out.append(row)
    return sorted(out, key=lambda item: (item["framework"], item["domain"], item["evidence_artifact_type"], item["response_path"])), unmapped


def questionnaire_evidence_summary_rows(enriched_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    for item in enriched_rows:
        artifact = item["evidence_artifact_type"]
        record = grouped.setdefault(
            artifact,
            {
                "evidence_artifact_type": artifact,
                "requirement_count": 0,
                "frameworks": set(),
                "domains": set(),
                "evidence_expectations": [],
            },
        )
        record["requirement_count"] += 1
        record["frameworks"].add(item["framework"])
        record["domains"].add(item["domain"])
        if item.get("evidence_expectation"):
            record["evidence_expectations"].append(item["evidence_expectation"])
    out: list[dict[str, Any]] = []
    for record in grouped.values():
        expectations = count_values(({"value": value} for value in record["evidence_expectations"]), "value")
        top_expectation = next(iter(expectations.keys()), "Provide evidence referenced in the XLSX.")
        out.append(
            {
                "evidence_artifact_type": record["evidence_artifact_type"],
                "requirement_count": record["requirement_count"],
                "frameworks": compact_join(sorted(record["frameworks"])),
                "domains": compact_join(sorted(record["domains"]), limit=6),
                "evidence_expectation": top_expectation,
            }
        )
    return sorted(out, key=lambda item: (-item["requirement_count"], item["evidence_artifact_type"]))


def shade_cell(cell: Any, fill: str) -> None:
    if OxmlElement is None or qn is None:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_widths(table: Any, widths: list[float]) -> None:
    table.autofit = False
    if OxmlElement is not None and qn is not None:
        width_dxa = [str(int(round(width * 1440))) for width in widths]
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(sum(int(value) for value in width_dxa)))

        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_ind.set(qn("w:w"), "120")

        tbl_layout = tbl_pr.find(qn("w:tblLayout"))
        if tbl_layout is None:
            tbl_layout = OxmlElement("w:tblLayout")
            tbl_pr.append(tbl_layout)
        tbl_layout.set(qn("w:type"), "fixed")

        cell_mar = tbl_pr.find(qn("w:tblCellMar"))
        if cell_mar is None:
            cell_mar = OxmlElement("w:tblCellMar")
            tbl_pr.append(cell_mar)
        for side, value in {"top": "80", "bottom": "80", "start": "120", "end": "120"}.items():
            node = cell_mar.find(qn(f"w:{side}"))
            if node is None:
                node = OxmlElement(f"w:{side}")
                cell_mar.append(node)
            node.set(qn("w:w"), value)
            node.set(qn("w:type"), "dxa")

        existing_grid = tbl.find(qn("w:tblGrid"))
        if existing_grid is not None:
            tbl.remove(existing_grid)
        tbl_grid = OxmlElement("w:tblGrid")
        for value in width_dxa:
            grid_col = OxmlElement("w:gridCol")
            grid_col.set(qn("w:w"), value)
            tbl_grid.append(grid_col)
        tbl.insert(1, tbl_grid)

    for row in table.rows:
        for index, width in enumerate(widths):
            if index >= len(row.cells):
                continue
            row.cells[index].width = Inches(width)
            if OxmlElement is not None and qn is not None:
                tc_pr = row.cells[index]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:type"), "dxa")
                tc_w.set(qn("w:w"), str(int(round(width * 1440))))


def set_cell_text(cell: Any, text: str, bold: bool = False, fill: str | None = None, font_size: int = 8) -> None:
    cell.text = text
    if fill:
        shade_cell(cell, fill)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            if fill in {"1F4E78", "0B2545"}:
                run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_docx_metadata_table(doc: Any, metadata: dict[str, Any]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_widths(table, [1.6, 7.0])
    set_cell_text(table.rows[0].cells[0], "Reference", True, "1F4E78")
    set_cell_text(table.rows[0].cells[1], "Value", True, "1F4E78")
    for key, value in metadata.items():
        cells = table.add_row().cells
        set_cell_text(cells[0], str(key), True)
        set_cell_text(cells[1], str(value))


def configure_docx_styles(doc: Any) -> None:
    styles = doc.styles
    for style_name, size, color in [
        ("Normal", 9, RGBColor(0, 0, 0)),
        ("Title", 22, RGBColor(31, 78, 121)),
        ("Heading 1", 14, RGBColor(46, 116, 181)),
        ("Heading 2", 11, RGBColor(31, 77, 120)),
        ("Heading 3", 10, RGBColor(31, 77, 120)),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        if hasattr(style, "paragraph_format"):
            style.paragraph_format.space_before = Pt(0 if style_name in {"Normal", "Title"} else 6)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.15


def add_body_text(doc: Any, text: str, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.italic = italic


def add_hidden_marker(doc: Any, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(1)
    run.font.hidden = True
    run.font.color.rgb = RGBColor(255, 255, 255)


def add_section_table(doc: Any, headers: list[str], rows_for_table: list[list[Any]], widths: list[float], font_size: int = 8) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, "1F4E78", font_size=font_size)
    for row_values in rows_for_table:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            set_cell_text(cells[index], "" if value is None else str(value), font_size=font_size)
    set_table_widths(table, widths)
    doc.add_paragraph("")
    return table


def add_blank_response_rows(count: int, code_prefix: str, prompts: list[str]) -> list[list[str]]:
    rows_for_table: list[list[str]] = []
    for index, prompt in enumerate(prompts, start=1):
        rows_for_table.append([f"{code_prefix}-{index:02d}", prompt, "", "", ""])
    while len(rows_for_table) < count:
        index = len(rows_for_table) + 1
        rows_for_table.append([f"{code_prefix}-{index:02d}", "", "", "", ""])
    return rows_for_table


def export_docx(conn: sqlite3.Connection, output_dir: Path, db_path: Path, xlsx_path: Path) -> Path:
    if Document is None:
        raise RuntimeError("python-docx is required to export DOCX")
    output_dir.mkdir(parents=True, exist_ok=True)
    dated = output_dir / f"regulatory_vendor_questionnaire_template_{today_slug()}.docx"
    latest = output_dir / "regulatory_vendor_questionnaire_template_latest.docx"
    requirement_count = conn.execute("SELECT COUNT(*) FROM requirements").fetchone()[0]
    source_count = conn.execute("SELECT COUNT(*) FROM sources").fetchone()[0]
    db_hash = sha256_file(db_path) if db_path.exists() else ""
    xlsx_hash = sha256_file(xlsx_path) if xlsx_path.exists() else ""
    enriched_rows = enriched_requirement_rows(conn)
    issues = source_issue_rows(conn)
    coverage_rows, unmapped_rows = questionnaire_coverage_rows(enriched_rows)
    if unmapped_rows:
        domains = compact_join(sorted({str(row.get("domain") or "Unknown") for row in unmapped_rows}))
        raise RuntimeError(f"DOCX questionnaire coverage has unmapped requirement domains: {domains}")
    mapped_count = sum(int(row["requirement_count"]) for row in coverage_rows)
    if mapped_count != requirement_count:
        raise RuntimeError(f"DOCX questionnaire coverage mismatch: mapped={mapped_count}, requirements={requirement_count}")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    configure_docx_styles(doc)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Vendor Regulatory Scoping Questionnaire")
    run.font.color.rgb = RGBColor(31, 78, 121)
    add_body_text(
        doc,
        "Concise vendor-facing questionnaire generated from the approved regulatory XLSX/index. "
        "The fillable sections collect the scope, evidence, approvals, exceptions, and attestation needed for a reviewer to confirm regulatory coverage.",
    )

    add_docx_metadata_table(
        doc,
        {
            "Generated at": now_iso(),
            "XLSX reference": rel(xlsx_path),
            "XLSX SHA-256": xlsx_hash,
            "Database": rel(db_path),
            "Database SHA-256": db_hash,
            "Requirement count": requirement_count,
            "Source count": source_count,
            "Known source issues": len(issues),
            "Mapped requirements total": mapped_count,
            "Unmapped requirements": 0,
        },
    )

    doc.add_heading("Completion Instructions", level=1)
    add_section_table(
        doc,
        ["Step", "Vendor action", "Required result"],
        [
            ["1", "Complete every fillable table in this DOCX.", "Reviewer can confirm scope, evidence, approvals, and exceptions."],
            ["2", "Use the XLSX Requirements and Vendor Questions tabs as the full reference.", "All requirement IDs remain traceable to source text and evidence expectations."],
            ["3", "Reference evidence by attachment name, secure portal link, GRC record, report section, or policy identifier.", "No secrets, credentials, customer data, or regulated records are inserted into this document."],
            ["4", "Record every non-applicable item, gap, or exception.", "Reviewer can assess residual regulatory exposure before test approval."],
            ["5", "Complete the Vendor Attestation and Sign-Off section.", "Authorized vendor representative confirms completeness and accuracy."],
        ],
        [0.55, 4.25, 4.65],
    )

    doc.add_heading("Regulatory Pack Summary", level=1)
    summary_rows = []
    for framework in sorted({row["framework"] for row in enriched_rows}, key=lambda value: (value == "SCOPE_TEMPLATE", value)):
        framework_items = [row for row in enriched_rows if row["framework"] == framework]
        summary_rows.append(
            [
                framework,
                len(framework_items),
                sum(1 for row in framework_items if row["priority"] in {"Critical", "High"}),
                sum(1 for row in framework_items if row["quality_flag"] != "OK"),
                compact_join(count_values(framework_items, "domain").keys(), limit=4),
            ]
        )
    add_section_table(
        doc,
        ["Framework", "Requirements", "High priority", "Review flags", "Primary domains"],
        summary_rows,
        [1.35, 1.05, 1.05, 1.05, 4.95],
    )

    doc.add_heading("Vendor and Test Profile", level=1)
    add_section_table(
        doc,
        ["Field", "Vendor response", "Reviewer use"],
        [
            ["Vendor legal name", "", "Confirm contracting and evidence ownership."],
            ["Primary vendor contact and role", "", "Route clarifications and final attestation."],
            ["Client/financial entity", "", "Tie scope to the regulated entity or business service."],
            ["Engagement/test name", "", "Align with SOW, rules of engagement, and report title."],
            ["Planned test dates and report due date", "", "Confirm regulatory timing and remediation windows."],
            ["Jurisdictions/regulators considered", "", "Confirm DORA, DORA TLPT, FRB, FFIEC, HKMA iCAST, PRA CBEST, and supplied scope applicability."],
            ["Evidence repository or secure portal", "", "Record where supporting artifacts will be provided."],
            ["Primary reviewer or approval board", "", "Identify who will judge regulatory completeness."],
        ],
        [2.15, 4.55, 2.75],
    )

    doc.add_heading("Regulatory Applicability", level=1)
    applicability_rows = []
    for framework in sorted({row["framework"] for row in enriched_rows}, key=lambda value: (value == "SCOPE_TEMPLATE", value)):
        framework_items = [row for row in enriched_rows if row["framework"] == framework]
        applicability_rows.append(
            [
                framework,
                len(framework_items),
                compact_join(count_values(framework_items, "domain").keys(), limit=4),
                "Applies / Does not apply / Review",
                "",
                "",
            ]
        )
    add_section_table(
        doc,
        ["Framework", "Reqs", "Coverage areas", "Applicability", "Vendor rationale", "Evidence reference"],
        applicability_rows,
        [1.05, 0.55, 2.45, 1.55, 2.3, 1.55],
    )

    doc.add_heading("Regulatory Scoping Questionnaire", level=1)
    add_body_text(
        doc,
        "Complete each response path below. The Regulatory Coverage Matrix maps every requirement row to one of these response paths.",
        italic=True,
    )
    scoping_rows = [
        [row["code"], row["coverage"], row["prompt"], "", ""]
        for row in QUESTIONNAIRE_CORE_ROWS
    ]
    add_section_table(
        doc,
        ["Path", "Regulatory coverage", "Vendor must provide", "Vendor response / scope detail", "Evidence reference / owner / status"],
        scoping_rows,
        [0.75, 2.0, 3.0, 2.0, 1.7],
        font_size=7,
    )

    doc.add_heading("Evidence Package Checklist", level=1)
    evidence_rows = []
    for item in questionnaire_evidence_summary_rows(enriched_rows):
        evidence_rows.append(
            [
                item["evidence_artifact_type"],
                item["requirement_count"],
                item["frameworks"],
                item["domains"],
                item["evidence_expectation"],
                "",
            ]
        )
    add_section_table(
        doc,
        ["Evidence type", "Reqs", "Frameworks", "Domains", "Expected evidence", "Vendor evidence reference / status"],
        evidence_rows,
        [1.65, 0.55, 1.25, 1.8, 2.65, 1.55],
        font_size=7,
    )

    doc.add_heading("Exception and Remediation Register", level=1)
    add_section_table(
        doc,
        ["Item", "Requirement path or evidence area", "Exception / gap", "Compensating control", "Owner / target date / status"],
        add_blank_response_rows(
            8,
            "EXC",
            [
                "Any out-of-scope system, geography, data class, service, or provider that may affect regulatory coverage.",
                "Any missing evidence artifact or delayed deliverable.",
                "Any test limitation, safety constraint, credential limitation, or access constraint.",
                "Any unconfirmed regulator/client approval or management sign-off.",
            ],
        ),
        [0.55, 2.25, 2.2, 2.2, 2.25],
    )

    doc.add_heading("Source Review Acknowledgement", level=1)
    if issues:
        source_rows = [
            [
                issue["source_id"],
                issue["framework"],
                issue["issue_type"],
                issue["recommended_action"],
                "",
            ]
            for issue in issues
        ]
    else:
        source_rows = [["No known source issues", "", "No issue recorded", "Confirm the XLSX Source Issues tab was reviewed.", ""]]
    add_section_table(
        doc,
        ["Source", "Framework", "Issue", "Required acknowledgement", "Vendor/reviewer response"],
        source_rows,
        [2.05, 1.05, 1.45, 3.0, 1.9],
        font_size=7,
    )

    doc.add_heading("Regulatory Coverage Matrix", level=1)
    add_body_text(
        doc,
        f"Mapped requirements total: {mapped_count}. Unmapped requirements: 0. "
        "This matrix proves coverage for every requirement row without repeating full regulatory text in Word.",
        italic=True,
    )
    coverage_table_rows = [
        [
            row["framework"],
            row["domain"],
            row["evidence_artifact_type"],
            row["requirement_count"],
            row["response_path"],
            row["requirement_ids"],
        ]
        for row in coverage_rows
    ]
    add_section_table(
        doc,
        ["Framework", "Domain", "Evidence type", "Reqs", "Response path", "Requirement ID sample"],
        coverage_table_rows,
        [1.0, 1.75, 1.85, 0.45, 0.95, 3.45],
        font_size=6,
    )

    doc.add_heading("Source Register", level=1)
    source_rows = []
    for source in rows(conn, "SELECT framework, title, authority, status, local_path, sha256 FROM sources ORDER BY framework, title"):
        source_rows.append(
            [
                source["framework"],
                source["title"],
                source["authority"],
                source["status"],
                source["local_path"],
                (source["sha256"] or "")[:16],
            ]
        )
    add_section_table(
        doc,
        ["Framework", "Title", "Authority", "Status", "Local path", "SHA-256 prefix"],
        source_rows,
        [0.9, 2.15, 1.6, 0.85, 3.15, 0.8],
        font_size=7,
    )

    doc.add_heading("Vendor Attestation and Sign-Off", level=1)
    add_body_text(
        doc,
        "The vendor must complete this section before the questionnaire is accepted for regulatory review.",
        italic=True,
    )
    add_hidden_marker(doc, "ATTESTATION_REQUIRED")
    add_section_table(
        doc,
        ["Attestation item", "Vendor confirmation / notes"],
        [
            ["I confirm the scope responses are complete and accurate for the proposed test.", ""],
            ["I confirm evidence references identify current, reviewable artifacts and do not include secrets or regulated records in this document.", ""],
            ["I confirm all exceptions, exclusions, unresolved source issues, and remediation dependencies are recorded above.", ""],
            ["I confirm the reviewer can use this DOCX with the referenced XLSX to assess coverage for all regulatory requirements.", ""],
            ["Authorized vendor representative name, title, signature, and date", ""],
            ["Client/reviewer acceptance name, title, signature, and date", ""],
        ],
        [3.45, 6.0],
    )

    add_hidden_marker(doc, f"XLSX_VENDOR_WORKLIST_REQUIRED: {xlsx_path.name}")
    add_hidden_marker(doc, f"MAPPED_REQUIREMENTS_TOTAL:{mapped_count}")
    add_hidden_marker(doc, "UNMAPPED_REQUIREMENTS:0")

    doc.save(dated)
    shutil.copyfile(dated, latest)
    return dated


def prepare_artifact_run(conn: sqlite3.Connection, artifact_type: str, path: Path, notes: str = "") -> None:
    requirement_count = conn.execute("SELECT COUNT(*) FROM requirements").fetchone()[0]
    source_count = conn.execute("SELECT COUNT(*) FROM sources").fetchone()[0]
    run_id = stable_id("RUN", artifact_type, str(path))
    # Artifact hashes are intentionally omitted from index.db. Writing the
    # DOCX hash into the database would change the database hash that the DOCX
    # itself records, creating a circular reference.
    conn.execute(
        """
        INSERT OR REPLACE INTO artifact_runs (
            run_id, artifact_type, path, sha256, generated_at,
            requirement_count, source_count, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (run_id, artifact_type, rel(path), None, now_iso(), requirement_count, source_count, notes),
    )


def export_artifacts(db_path: Path = DEFAULT_DB, output_dir: Path = DEFAULT_OUTPUT_DIR) -> dict[str, Any]:
    ensure_dependencies("export")
    if not db_path.is_file():
        raise FileNotFoundError(f"Database not found: {db_path}")
    conn = sqlite3.connect(str(db_path))
    try:
        xlsx_path = output_dir / f"regulatory_requirements_reference_{today_slug()}.xlsx"
        docx_path = output_dir / f"regulatory_vendor_questionnaire_template_{today_slug()}.docx"
        conn.execute("DELETE FROM artifact_runs")
        prepare_artifact_run(conn, "xlsx", xlsx_path, "Regulatory reference workbook")
        prepare_artifact_run(conn, "docx", docx_path, "Vendor-fillable requirements questionnaire")
        conn.commit()
        xlsx_path = export_xlsx(conn, output_dir, db_path)
        docx_path = export_docx(conn, output_dir, db_path, xlsx_path)
        summary = validate_outputs(db_path, xlsx_path, docx_path)
        summary.update(
            {
                "xlsx_path": rel(xlsx_path),
                "xlsx_latest_path": rel(output_dir / "regulatory_requirements_reference_latest.xlsx"),
                "docx_path": rel(docx_path),
                "docx_latest_path": rel(output_dir / "regulatory_vendor_questionnaire_template_latest.docx"),
            }
        )
        return summary
    finally:
        conn.close()


def clean_generated(output_dir: Path = DEFAULT_OUTPUT_DIR) -> dict[str, Any]:
    keep_names = {
        "regulatory_requirements_reference_latest.xlsx",
        "regulatory_vendor_questionnaire_template_latest.docx",
    }
    removed: list[str] = []
    targets: list[Path] = [
        output_dir / "index.db",
        output_dir / "download-log.json",
    ]
    targets.extend(output_dir.glob("regulatory_requirements_reference_*.xlsx"))
    targets.extend(output_dir.glob("regulatory_vendor_questionnaire_template_*.docx"))
    for target in sorted(set(targets)):
        if target.name in keep_names or not target.exists():
            continue
        target.unlink()
        removed.append(rel(target))
    for directory_name in ("docx_render_check", "xlsx_render_check"):
        directory = output_dir / directory_name
        if directory.exists():
            shutil.rmtree(directory)
            removed.append(rel(directory))
    return {"removed": removed, "kept": sorted(keep_names)}


def validate_outputs(db_path: Path, xlsx_path: Path, docx_path: Path) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to validate XLSX")
    db_hash = sha256_file(db_path) if db_path.exists() else ""
    conn = sqlite3.connect(str(db_path))
    try:
        db_requirements = conn.execute("SELECT COUNT(*) FROM requirements").fetchone()[0]
        frameworks = [row[0] for row in conn.execute("SELECT framework FROM requirements GROUP BY framework HAVING COUNT(*) > 0")]
        domains = [row[0] for row in conn.execute("SELECT domain FROM requirements GROUP BY domain HAVING COUNT(*) > 0")]
        enriched_rows = enriched_requirement_rows(conn)
    finally:
        conn.close()
    coverage_rows, unmapped_rows = questionnaire_coverage_rows(enriched_rows)
    mapped_count = sum(int(row["requirement_count"]) for row in coverage_rows)
    wb = load_workbook(xlsx_path, data_only=False, read_only=False)
    try:
        xlsx_requirements = max(0, wb["Requirements"].max_row - 1)
        required_sheets = {
            "Dashboard",
            "Read Me",
            "Sources",
            "Requirements",
            "Vendor Questions",
            "Vendor Worklist",
            "Vendor - DORA",
            "Vendor - US Bank",
            "Vendor - HKMA CBEST",
            "Vendor - Scope",
            "Evidence Matrix",
            "Evidence Checklist",
            "Service Mapping",
            "QA Flags",
            "Source Issues",
            "Data Dictionary",
            "Change Log",
        }
        missing_sheets = sorted(required_sheets.difference(wb.sheetnames))
        vendor_validations = len(wb["Vendor Questions"].data_validations.dataValidation)
        vendor_headers = [cell.value for cell in wb["Vendor Questions"][1]]
    finally:
        wb.close()
    docx_size = docx_path.stat().st_size if docx_path.exists() else 0
    if db_requirements != xlsx_requirements:
        raise RuntimeError(f"XLSX requirement count mismatch: db={db_requirements}, xlsx={xlsx_requirements}")
    if missing_sheets:
        raise RuntimeError(f"XLSX is missing required sheets: {', '.join(missing_sheets)}")
    for header in ["response_required", "applicability", "implementation_status", "evidence_status", "completion_status"]:
        if header not in vendor_headers:
            raise RuntimeError(f"Vendor Questions sheet is missing required vendor field: {header}")
    if vendor_validations < 5:
        raise RuntimeError("Vendor Questions sheet is missing expected dropdown validations")
    if not docx_path.exists() or docx_size < 1000:
        raise RuntimeError("DOCX output is missing or unexpectedly small")
    with zipfile.ZipFile(docx_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    document_text = html.unescape(re.sub(r"<[^>]+>", "", document_xml))
    if xlsx_path.name not in document_xml or db_hash not in document_xml:
        raise RuntimeError("DOCX metadata does not reference the exported XLSX and database hash")
    if unmapped_rows:
        domains_missing = compact_join(sorted({str(row.get("domain") or "Unknown") for row in unmapped_rows}))
        raise RuntimeError(f"DOCX coverage matrix has unmapped requirement domains: {domains_missing}")
    if mapped_count != db_requirements:
        raise RuntimeError(f"DOCX coverage matrix count mismatch: mapped={mapped_count}, db={db_requirements}")
    required_docx_sections = [
        "Vendor and Test Profile",
        "Regulatory Applicability",
        "Regulatory Scoping Questionnaire",
        "Evidence Package Checklist",
        "Exception and Remediation Register",
        "Source Review Acknowledgement",
        "Regulatory Coverage Matrix",
        "Vendor Attestation and Sign-Off",
        "ATTESTATION_REQUIRED",
        f"MAPPED_REQUIREMENTS_TOTAL:{db_requirements}",
        "UNMAPPED_REQUIREMENTS:0",
    ]
    for section_name in required_docx_sections:
        if section_name not in document_text:
            raise RuntimeError(f"DOCX questionnaire is missing required section or marker: {section_name}")
    for framework in frameworks:
        if framework not in document_text:
            raise RuntimeError(f"DOCX coverage matrix is missing framework: {framework}")
    for domain in domains:
        if domain not in document_text:
            raise RuntimeError(f"DOCX coverage matrix is missing domain: {domain}")
    return {
        "database_requirements": db_requirements,
        "xlsx_requirements": xlsx_requirements,
        "frameworks": frameworks,
        "docx_mapped_requirements": mapped_count,
        "docx_size_bytes": docx_size,
        "validated_at": now_iso(),
    }


def print_json(payload: dict[str, Any] | list[dict[str, Any]]) -> None:
    print(json.dumps(payload, indent=2, sort_keys=True))


def command_download(args: argparse.Namespace) -> int:
    sources = load_sources(args.manifest)
    results = download_sources(sources, args.output_dir, force=args.force, offline=args.offline, refresh=args.refresh)
    print_json({"download_log": rel(download_log_path(args.output_dir)), "results": results})
    failed = [item for item in results if item["status"] == "failed"]
    return 1 if failed and args.fail_on_download_error else 0


def command_refresh_sources(args: argparse.Namespace) -> int:
    sources = load_sources(args.manifest)
    results = download_sources(sources, args.output_dir, force=False, offline=False, refresh=True)
    print_json({"download_log": rel(download_log_path(args.output_dir)), "results": results})
    failed = [item for item in results if item["status"] == "failed"]
    return 1 if failed and args.fail_on_download_error else 0


def command_index(args: argparse.Namespace) -> int:
    sources = load_sources(args.manifest)
    summary = build_index(sources, args.db, args.output_dir)
    print_json(summary)
    return 0


def command_export(args: argparse.Namespace) -> int:
    summary = export_artifacts(args.db, args.output_dir)
    print_json(summary)
    return 0


def command_build_all(args: argparse.Namespace) -> int:
    sources = load_sources(args.manifest)
    download_results = download_sources(
        sources,
        args.output_dir,
        force=args.force,
        offline=args.offline,
        refresh=args.refresh,
    )
    index_summary = build_index(sources, args.db, args.output_dir)
    export_summary = export_artifacts(args.db, args.output_dir)
    failed = [item for item in download_results if item["status"] == "failed"]
    print_json(
        {
            "downloads": {
                "total": len(download_results),
                "failed": len(failed),
                "log": rel(download_log_path(args.output_dir)),
            },
            "index": index_summary,
            "exports": export_summary,
        }
    )
    return 1 if failed and args.fail_on_download_error else 0


def command_clean(args: argparse.Namespace) -> int:
    print_json(clean_generated(args.output_dir))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="regulatory-requirements",
        description=(
            "Download official regulatory sources, index clause-level requirements, "
            f"and export XLSX/DOCX artifacts for {REGULATORY_FOCUS}"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=textwrap.dedent(
            f"""
            Examples:
              python -m regulatory_requirements build-all --refresh
              python -m regulatory_requirements refresh-sources
              python -m regulatory_requirements build-all --offline
              python -m regulatory_requirements export --db {DEFAULT_DB}
              python -m regulatory_requirements clean
            """
        ),
    )
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST, help="Source manifest JSON path")
    common.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="Generated output directory")
    common.add_argument("--db", type=Path, default=DEFAULT_DB, help="SQLite index database path")

    sub = parser.add_subparsers(dest="command", required=True)

    download = sub.add_parser("download", parents=[common], help="Download official sources and write provenance log")
    download.add_argument("--force", action="store_true", help="Re-download even when a local archive exists")
    download.add_argument("--refresh", action="store_true", help="Fetch official sources even when cached and update only if changed")
    download.add_argument("--offline", action="store_true", help="Skip network and report existing local files")
    download.add_argument("--fail-on-download-error", action="store_true", help="Return non-zero if any download fails")
    download.set_defaults(func=command_download)

    refresh = sub.add_parser("refresh-sources", parents=[common], help="Check official regulator sources and update local archives when changed")
    refresh.add_argument("--fail-on-download-error", action="store_true", help="Return non-zero if any source refresh fails")
    refresh.set_defaults(func=command_refresh_sources)

    index = sub.add_parser("index", parents=[common], help="Build index.db from local source files")
    index.set_defaults(func=command_index)

    export = sub.add_parser("export", parents=[common], help="Export XLSX and DOCX from index.db")
    export.set_defaults(func=command_export)

    build_all = sub.add_parser("build-all", parents=[common], help="Run download, index, and export")
    build_all.add_argument("--force", action="store_true", help="Re-download even when a local archive exists")
    build_all.add_argument("--refresh", action="store_true", help="Check official regulator sources and update archives before indexing")
    build_all.add_argument("--offline", action="store_true", help="Skip network and use local/cached files")
    build_all.add_argument("--fail-on-download-error", action="store_true", help="Return non-zero if any download fails")
    build_all.set_defaults(func=command_build_all)

    clean = sub.add_parser("clean", parents=[common], help="Remove local generated files while keeping approved latest XLSX/DOCX outputs")
    clean.set_defaults(func=command_clean)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return int(args.func(args))
    except Exception as exc:  # noqa: BLE001 - CLI should return clear failures.
        print(f"error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
